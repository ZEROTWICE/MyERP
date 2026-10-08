# -*- coding: utf-8 -*-
"""t1（B12-02）验证：/production_records/import 的 notes KeyError 修复。

全程只用 app.db 副本（scripts/_test_bootstrap.make_app()），真实 app.db 只读；
`_t1_` 前缀标记本脚本自建的测试实体，避免与既有数据混淆。

判据（与任务 Acceptance 一一对应）：
  A1 parse_production_record_data() 每个 dict 都含 notes 键，缺列时 None，不抛错
  A2 修后：合法行 ProductionRecord Δ=+1，且 process_id 指向命中价格、价格快照可读
  A3 阴对照：同一夹具下把解析结果里 'notes' 去掉（等价修前状态）⇒ records_created=0
     且错误串含 "处理记录时出错: 'notes'"
  A4 行级容错不回归：坏行只记「第 N 行：…」并跳过，响应不整体 500
  A5 无 schema 变更；真库 app.db SHA256 不变

用法：python -B scripts/_sandbox_compat.py test-reports-2026-10/reconcile-2026-10-08/probe_t1_notes_keyerror.py
"""
import hashlib
import io
import json
import os
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'scripts'))

from _test_bootstrap import ensure_role_users, login_as, make_app  # noqa: E402

REAL_DB = ROOT / 'app.db'
REAL_SHA_EXPECT = 'F5DA2306BC31CBAB098AAD3774016D320A9B9AA3546F93418196AE9900E0F065'

EMP_TEST = '_t1_EMP9001'
PROC_TEST = '_T1P9001'
PROC_MISSING = '_T1P9999'
BAD_ROW_MSG_PREFIX = '员工工号'


def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest().upper()


def build_fixture_bytes():
    """4 列表头（与 create_production_record_template 一致）+ 1 条合法行 + 2 条坏行。"""
    from datetime import date, timedelta

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = '生产记录'
    ws.append(['日期*', '员工工号*', '工序编号*', '数量*'])
    # 合法行：命中 EMP_TEST / PROC_TEST
    ws.append([date.today() + timedelta(days=1), EMP_TEST, PROC_TEST, 7])
    # 坏行 1：员工不存在
    ws.append([date.today() + timedelta(days=1), '_t1_EMP0000', PROC_TEST, 3])
    # 坏行 2：工序无价格
    ws.append([date.today() + timedelta(days=1), EMP_TEST, PROC_MISSING, 3])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def seed(app):
    """自建测试实体（都在副本上）。"""
    from datetime import date, datetime, timedelta

    from app import db
    from app.models import Employee, ProcessPrice, SerialNumber

    with app.app_context():
        emp = Employee.query.filter_by(employee_id=EMP_TEST).first()
        if emp is None:
            emp = Employee(global_sn=SerialNumber.get_next_number(), employee_id=EMP_TEST,
                           name='_t1_测试员工', position='普通员工',
                           hire_date=date.today())
            db.session.add(emp)
            db.session.flush()
        effective = datetime.combine(date.today() + timedelta(days=1), datetime.min.time())
        pp = ProcessPrice.query.filter_by(process_code=PROC_TEST).first()
        if pp is None:
            pp = ProcessPrice(global_sn=SerialNumber.get_next_number(), process_code=PROC_TEST,
                              process_name='_t1_测试工序', price=1.25, version=1,
                              effective_date=effective)
            db.session.add(pp)
            db.session.flush()
        db.session.commit()
        return emp.id, pp.id


def post_import(app, client, monkeypatch_strip_notes=False):
    """发一次导入请求，返回 (status_code, json)。"""
    from app.models import ProductionRecord
    from app.utils.excel_generator import ExcelGenerator

    if monkeypatch_strip_notes:
        original = ExcelGenerator.parse_production_record_data

        def stripped(file_path):
            return [{k: v for k, v in row.items() if k != 'notes'}
                    for row in original(file_path)]

        ExcelGenerator.parse_production_record_data = staticmethod(stripped)

    with app.app_context():
        before = ProductionRecord.query.count()

    data = {'file': (build_fixture_bytes(), 'fixture_t1.xlsx')}
    resp = client.post('/production_records/import', data=data,
                       content_type='multipart/form-data')
    if monkeypatch_strip_notes:
        ExcelGenerator.parse_production_record_data = original

    with app.app_context():
        after = ProductionRecord.query.count()

    try:
        payload = resp.get_json(silent=True)
    except Exception:  # noqa: BLE001
        payload = None
    return resp.status_code, payload, after - before


def main():
    real_sha_before = sha256(REAL_DB)
    result = {'real_db_sha256_before': real_sha_before,
              'real_db_sha256_expected': REAL_SHA_EXPECT}

    # ---- A1：解析器直接验证（含 4 列与 5 列两种夹具） ----
    import tempfile

    from openpyxl import Workbook
    from app.utils.excel_generator import ExcelGenerator

    tmpdir = tempfile.mkdtemp(prefix='t1_fixture_')
    p4 = os.path.join(tmpdir, 'four_col.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.append(['日期*', '员工工号*', '工序编号*', '数量*'])
    ws.append(['2026-10-08', EMP_TEST, PROC_TEST, 5])
    wb.save(p4)
    p5 = os.path.join(tmpdir, 'five_col.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.append(['日期*', '员工工号*', '工序编号*', '数量*', '备注'])
    ws.append(['2026-10-08', EMP_TEST, PROC_TEST, 5, '  夜班补录  '])
    wb.save(p5)
    p5b = os.path.join(tmpdir, 'five_col_blank.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.append(['日期*', '员工工号*', '工序编号*', '数量*', '备注'])
    ws.append(['2026-10-08', EMP_TEST, PROC_TEST, 5, None])
    wb.save(p5b)

    rows4 = ExcelGenerator.parse_production_record_data(p4)
    rows5 = ExcelGenerator.parse_production_record_data(p5)
    rows5b = ExcelGenerator.parse_production_record_data(p5b)
    result['A1_parse'] = {
        'four_col_keys': sorted(rows4[0].keys()),
        'four_col_notes': rows4[0]['notes'],
        'five_col_notes': rows5[0]['notes'],
        'five_col_blank_notes': rows5b[0]['notes'],
        'all_dicts_have_notes': all('notes' in r for r in rows4 + rows5 + rows5b),
        'no_exception': True,
    }

    # ---- 起应用（副本）----
    app, copy_path = make_app(fresh=True)
    result['db_copy_path'] = copy_path
    result['db_uri'] = app.config['SQLALCHEMY_DATABASE_URI']

    # 沙箱下 uploads/temp 被 ACL 拒写（仅环境层限制，与本任务代码无关）。
    # 端点只读 app.config['TEMP_FOLDER']，因此把它指到工作区内可写目录 —— 端点代码零改动。
    import tempfile as _tf
    scratch = ROOT / '.analysis-scratch' / 't1_temp'
    scratch.mkdir(parents=True, exist_ok=True)
    app.config['TEMP_FOLDER'] = str(scratch)
    result['temp_folder_override'] = str(scratch)
    result['temp_folder_default'] = str(ROOT / 'uploads' / 'temp')

    mappings, pw = ensure_role_users(app)
    emp_id, pp_id = seed(app)
    result['seeded'] = {'employee_pk': emp_id, 'process_price_pk': pp_id}

    client = app.test_client()
    login_as(client, mappings['admin'], pw)

    # ---- A3：阴对照（先跑，模拟修前：dict 中无 notes）----
    status_neg, payload_neg, delta_neg = post_import(app, client, monkeypatch_strip_notes=True)
    result['A3_negative_control'] = {
        'http_status': status_neg,
        'records_created': delta_neg,
        'message': (payload_neg or {}).get('message'),
    }

    # ---- A2 + A4：修后同一夹具 ----
    status_pos, payload_pos, delta_pos = post_import(app, client, monkeypatch_strip_notes=False)
    msg = (payload_pos or {}).get('message') or ''
    result['A2_A4_after_fix'] = {
        'http_status': status_pos,
        'records_created': delta_pos,
        'message': msg,
    }

    # 命中行的字段核对
    from app.models import ProcessPrice, ProductionRecord
    with app.app_context():
        rec = (ProductionRecord.query
               .filter_by(employee_id=emp_id)
               .order_by(ProductionRecord.id.desc()).first())
        pp = ProcessPrice.query.get(pp_id)
        result['A2_hit_row'] = {
            'record_id': rec.id if rec else None,
            'process_id': rec.process_id if rec else None,
            'expected_process_id': pp_id,
            'process_id_matches_process_price': bool(rec and rec.process_id == pp_id),
            'quantity': rec.quantity if rec else None,
            'notes': rec.notes if rec else 'MISSING',
            'process_price_value': float(pp.price) if pp else None,
        }
        # 行级容错：坏行未入库
        result['A4_bad_rows'] = {
            'bad_employee_pk_rows': ProductionRecord.query.filter_by(employee_id=None).count(),
            'bad_employee_not_created': (ProductionRecord.query
                                         .filter(ProductionRecord.quantity == 3)
                                         .count()) == 0,
            'has_line_prefixed_error': ('员工工号' in msg) or ('未找到工序' in msg),
            'not_500': status_pos == 200,
        }

    # ---- A5：真库 SHA + schema 无变更 ----
    real_sha_after = sha256(REAL_DB)
    result['real_db_sha256_after'] = real_sha_after
    result['A5_real_db_unchanged'] = (real_sha_after == real_sha_before)

    # schema 变更检查：副本库（经 create_app 的启动自愈后）与真库的 schema 定义逐表比对
    import sqlite3
    schema_q = ("SELECT type, name, sql FROM sqlite_master "
                "WHERE name NOT LIKE 'sqlite_%' ORDER BY type, name")

    def schema_of(db_file):
        con = sqlite3.connect(db_file)
        try:
            return con.execute(schema_q).fetchall()
        finally:
            con.close()

    real_schema = schema_of(str(REAL_DB))
    copy_schema = schema_of(copy_path)
    diff = [r for r in copy_schema if r not in real_schema]
    result['A5_schema'] = {
        'real_objects': len(real_schema),
        'copy_objects': len(copy_schema),
        'copy_extra_objects': diff,
        'schema_identical': real_schema == copy_schema,
    }

    print(json.dumps(result, ensure_ascii=False, indent=1))
    out_path = pathlib.Path(__file__).with_suffix('.output.json')
    out_path.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding='utf-8')

    # 退出码即判据
    ok = (
        result['A1_parse']['all_dicts_have_notes']
        and result['A1_parse']['four_col_notes'] is None
        and result['A1_parse']['five_col_notes'] == '夜班补录'
        and result['A1_parse']['five_col_blank_notes'] is None
        and result['A3_negative_control']['records_created'] == 0
        and "处理记录时出错: 'notes'" in (result['A3_negative_control']['message'] or '')
        and result['A2_A4_after_fix']['records_created'] == 1
        and result['A2_hit_row']['process_id_matches_process_price']
        and result['A2_hit_row']['notes'] is None
        and result['A4_bad_rows']['not_500']
        and result['A4_bad_rows']['bad_employee_not_created']
        and result['A5_real_db_unchanged']
        and result['A5_schema']['schema_identical']
        and real_sha_after == REAL_SHA_EXPECT
    )
    print('RESULT:', 'OK' if ok else 'FAIL')
    return 0 if ok else 1


if __name__ == '__main__':
    raise SystemExit(main())
