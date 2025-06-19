#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
工价标准数据转换脚本
将合金钢辙叉计件工价标准.xlsx的数据转换并填入到process_price_template.xlsx
"""

import pandas as pd
import openpyxl
from openpyxl import load_workbook
import re
from pypinyin import lazy_pinyin, Style
from datetime import date


def chinese_to_pinyin_initial(text):
    """将中文转换为拼音首字母大写"""
    if not text:
        return text
    
    # 获取拼音首字母
    pinyin_list = lazy_pinyin(text, style=Style.FIRST_LETTER)
    return ''.join([p.upper() for p in pinyin_list])


def clean_code_string(text):
    """清理字符串中的符号，保留字母数字和中文"""
    if not text:
        return text
    
    # 移除常见符号
    cleaned = re.sub(r'[^\w\u4e00-\u9fff]', '', str(text))
    
    # 将中文转换为拼音首字母
    result = ''
    for char in cleaned:
        if '\u4e00' <= char <= '\u9fff':  # 是中文字符
            result += chinese_to_pinyin_initial(char)
        else:
            result += char
    
    return result


def split_by_slash(text):
    """根据斜杠分割文本"""
    if not text or '/' not in str(text):
        return [text] if text else ['']
    return [item.strip() for item in str(text).split('/')]


def get_cell_value(cell, formula_cell=None):
    """获取单元格的值，如果是公式则返回计算结果"""
    if cell is None:
        return None
    
    # 如果data_only模式的值存在且不为None，使用它
    if cell.value is not None:
        return cell.value
    
    # 如果data_only模式返回None，但有对应的公式单元格，尝试使用公式字符串
    if formula_cell is not None and hasattr(formula_cell, 'data_type') and formula_cell.data_type == 'f':
        # 这是一个公式，但无法计算结果，返回公式字符串作为标识
        return formula_cell.value
    
    return cell.value


def generate_process_code(model, drawing_num, seq_num):
    """生成工序编号"""
    # 清理型号和图号
    clean_model = clean_code_string(model) if model else ''
    clean_drawing = clean_code_string(drawing_num) if drawing_num else ''
    
    # 处理序号：如果seq_num是数字，格式化为3位；否则直接使用
    if seq_num and str(seq_num).strip() and str(seq_num).strip().isdigit():
        seq_str = f"{int(seq_num):03d}"
    elif seq_num and str(seq_num).strip():
        # 如果序号不是纯数字，直接使用（可能包含字母等）
        seq_str = clean_code_string(str(seq_num).strip())
    else:
        seq_str = "000"
    
    # 生成编号
    code = f"YHGH{clean_model}{clean_drawing}{seq_str}"
    return code


def process_source_data():
    """处理源数据文件"""
    print("正在读取源数据文件...")
    
    try:
        # 尝试以data_only模式加载源数据文件以获取公式计算结果
        # 如果公式结果为None，则回退到普通模式
        try:
            wb_source = load_workbook('合金钢辙叉计件工价标准.xlsx', data_only=True)
            wb_source_formula = load_workbook('合金钢辙叉计件工价标准.xlsx', data_only=False)
        except Exception as e:
            print(f"警告: 加载文件时出错: {e}")
            wb_source = load_workbook('合金钢辙叉计件工价标准.xlsx')
            wb_source_formula = wb_source
        
        # 跳过第一个工作表（00-工价目录）
        sheet_names = [name for name in wb_source.sheetnames if not name.startswith('00-')]
        
        print(f"找到 {len(sheet_names)} 个有效工作表")
        
        all_data = []
        error_sheets = []  # 存储出现错误的工作表信息
        
        for sheet_name in sheet_names:
            print(f"正在处理工作表: {sheet_name}")
            
            # 提取型号（工作表名称去掉序号前缀）
            model = re.sub(r'^\d+-', '', sheet_name)
            
            # 读取工作表数据
            ws = wb_source[sheet_name]
            
            # 找到数据开始行（通常包含表头）
            # 同时获取对应的公式工作表以备用
            ws_formula = wb_source_formula[sheet_name] if wb_source_formula != wb_source else None
            
            data_rows = []
            for row_idx, row in enumerate(ws.iter_rows()):
                # 处理每个单元格，确保公式被计算为结果值
                row_values = []
                formula_row = list(ws_formula.iter_rows())[row_idx] if ws_formula else None
                
                for col_idx, cell in enumerate(row):
                    formula_cell = formula_row[col_idx] if formula_row and col_idx < len(formula_row) else None
                    value = get_cell_value(cell, formula_cell)
                    row_values.append(value)
                
                if any(row_values):  # 跳过空行
                    data_rows.append(tuple(row_values))
            
            if not data_rows:
                continue
            
            # 寻找表头行（可能有多个表头行，找到最后一个完整的表头）
            header_row_idx = -1
            for i, row in enumerate(data_rows):
                if row and len(row) >= 6:
                    # 检查是否包含标准表头字段
                    row_str = [str(cell).strip() if cell else '' for cell in row]
                    if ('序号' in row_str and '型号' in row_str and '图号' in row_str and 
                        '部件名称' in row_str and '工序名称' in row_str and '单价' in row_str):
                        header_row_idx = i  # 记录最后一个找到的表头行
                        # 不立即break，继续寻找可能存在的后续表头行
            
            if header_row_idx == -1:
                print(f"  警告: 在工作表 {sheet_name} 中未找到表头行")
                # 记录错误的工作表
                error_sheets.append({
                    'sheet_name': sheet_name,
                    'error_type': '未找到表头行',
                    'worksheet': ws
                })
                continue
            
            # 获取表头
            headers = data_rows[header_row_idx]
            
            # 找到各列的索引（根据实际结构：序号、型号、图号、部件名称、工序名称、单价、备注）
            seq_col = -1
            model_col = -1
            drawing_col = -1
            part_name_col = -1
            process_name_col = -1
            price_col = -1
            remark_col = -1
            
            for i, header in enumerate(headers):
                if header and '序号' in str(header):
                    seq_col = i
                elif header and '型号' in str(header):
                    model_col = i
                elif header and '图号' in str(header):
                    drawing_col = i
                elif header and '部件名称' in str(header):
                    part_name_col = i
                elif header and '工序名称' in str(header):
                    process_name_col = i
                elif header and ('单价' in str(header) or '价格' in str(header)):
                    price_col = i
                elif header and '备注' in str(header):
                    remark_col = i
            
            if drawing_col == -1 or part_name_col == -1 or process_name_col == -1:
                print(f"  警告: 在工作表 {sheet_name} 中未找到必要的列")
                print(f"  表头: {headers}")
                # 记录错误的工作表
                error_sheets.append({
                    'sheet_name': sheet_name,
                    'error_type': '未找到必要的列',
                    'worksheet': ws
                })
                continue
            
            # 处理数据行
            sheet_data = []
            # 按图号分组存储工序数据
            normal_processes_by_drawing = {}  # 存储普通工序，按图号分组
            subtotal_processes_by_drawing = {}  # 存储小计工序，按图号分组
            current_part_processes_by_drawing = {}  # 存储当前部件的普通工序，按图号分组
            previous_part_name = ''  # 记录前一行的部件名称
            
            sequence = 1
            
            for i in range(header_row_idx + 1, len(data_rows)):
                row = data_rows[i]
                
                if not any(row):  # 跳过空行
                    continue
                
                # 获取各列数据
                seq_num = str(row[seq_col]) if seq_col != -1 and seq_col < len(row) and row[seq_col] else ''
                row_model = str(row[model_col]) if model_col != -1 and model_col < len(row) and row[model_col] else ''
                drawing_num = str(row[drawing_col]) if drawing_col < len(row) and row[drawing_col] else ''
                part_name = str(row[part_name_col]) if part_name_col < len(row) and row[part_name_col] else ''
                process_name = str(row[process_name_col]) if process_name_col < len(row) and row[process_name_col] else ''
                price = row[price_col] if price_col < len(row) and price_col != -1 and row[price_col] else 0
                remark = str(row[remark_col]) if remark_col < len(row) and remark_col != -1 and row[remark_col] else ''
                
                # 使用行内的型号（如果有），否则使用从工作表名称提取的型号
                actual_model = row_model if row_model else model
                
                # 跳过无效数据
                if not drawing_num and not part_name and not process_name:
                    continue
                
                # 处理图号分割
                drawing_nums = split_by_slash(drawing_num)
                
                # 工序名称不再分割，直接使用原始值
                process = process_name
                
                # 记录当前行的原始数据，用于小计和合计的包含工序统计
                current_row_processes_by_drawing = {}
                
                # 为每个图号创建记录
                for drawing in drawing_nums:
                    if not drawing and not process:
                        continue
                    
                    # 初始化图号分组字典
                    if drawing not in normal_processes_by_drawing:
                        normal_processes_by_drawing[drawing] = []
                    if drawing not in subtotal_processes_by_drawing:
                        subtotal_processes_by_drawing[drawing] = []
                    if drawing not in current_part_processes_by_drawing:
                        current_part_processes_by_drawing[drawing] = []
                    
                    # 生成工序编号，使用表中的序号列
                    process_code = generate_process_code(actual_model, drawing, seq_num)
                    
                    # 确定价格类型和包含工序
                    if part_name and '小计' in part_name:
                        price_type = 'subtotal'
                        # 小计包含与上一行部件名相同的当前部件组的普通工序（仅限相同图号）
                        included_processes = [p['process_code'] for p in current_part_processes_by_drawing[drawing]]
                    elif part_name and '合计' in part_name:
                        price_type = 'subtotal'
                        # 如果有小计工序，包含所有小计；如果没有小计，包含所有普通工序（仅限相同图号）
                        if subtotal_processes_by_drawing[drawing]:
                            included_processes = [p['process_code'] for p in subtotal_processes_by_drawing[drawing]]
                        else:
                            included_processes = [p['process_code'] for p in normal_processes_by_drawing[drawing]]
                    else:
                        price_type = 'normal'
                        included_processes = []
                    
                    process_data = {
                        'model': actual_model,
                        'drawing_num': drawing,
                        'part_name': part_name,
                        'process_name': process,
                        'process_code': process_code,
                        'price': price,
                        'price_type': price_type,
                        'included_processes': included_processes,
                        'remark': remark,
                        'sheet_name': sheet_name
                    }
                    
                    sheet_data.append(process_data)
                    all_data.append(process_data)
                    
                    # 记录当前图号的工序
                    current_row_processes_by_drawing[drawing] = process_data
                
                # 根据类型存储到相应列表（按图号分组）
                for drawing, process_data in current_row_processes_by_drawing.items():
                    if process_data['price_type'] == 'normal':
                        normal_processes_by_drawing[drawing].append(process_data)
                        # 跟踪当前部件的普通工序
                        if part_name == previous_part_name or previous_part_name == '':
                            # 如果是同一个部件或是第一行，添加到当前部件工序
                            current_part_processes_by_drawing[drawing].append(process_data)
                        else:
                            # 如果是新部件，重置当前部件工序列表
                            current_part_processes_by_drawing[drawing] = [process_data]
                    elif process_data['price_type'] == 'subtotal':
                        subtotal_processes_by_drawing[drawing].append(process_data)
                        # 小计后，清空当前部件工序（准备下一个部件）
                        if '小计' in part_name:
                            current_part_processes_by_drawing[drawing] = []
                
                # 更新前一行的部件名称（只对普通工序更新）
                if current_row_processes_by_drawing:
                    # 取第一个图号的工序类型作为参考
                    first_drawing = list(current_row_processes_by_drawing.keys())[0]
                    if current_row_processes_by_drawing[first_drawing]['price_type'] == 'normal':
                        previous_part_name = part_name
                
                sequence += 1
            
            print(f"  处理完成，共 {len(sheet_data)} 条记录")
        
        # 导出错误的工作表
        if error_sheets:
            export_error_sheets(error_sheets)
        
        return all_data
        
    except Exception as e:
        print(f"处理源数据时出错: {e}")
        return []


def export_error_sheets(error_sheets):
    """导出出现错误的工作表内容"""
    print(f"\n正在导出 {len(error_sheets)} 个错误工作表...")
    
    try:
        # 创建新的工作簿
        from openpyxl import Workbook
        wb_error = Workbook()
        
        # 删除默认工作表
        wb_error.remove(wb_error.active)
        
        # 创建错误信息汇总表
        ws_summary = wb_error.create_sheet("错误汇总")
        summary_headers = ['工作表名称', '错误类型', '原因说明']
        for col, header in enumerate(summary_headers, 1):
            ws_summary.cell(row=1, column=col, value=header)
        
        # 写入错误汇总信息
        for row_idx, error_info in enumerate(error_sheets, 2):
            ws_summary.cell(row=row_idx, column=1, value=error_info['sheet_name'])
            ws_summary.cell(row=row_idx, column=2, value=error_info['error_type'])
            if error_info['error_type'] == '未找到表头行':
                ws_summary.cell(row=row_idx, column=3, value='工作表中没有标准的表头行（序号、型号、图号、部件名称、工序名称、单价、备注）')
            elif error_info['error_type'] == '未找到必要的列':
                ws_summary.cell(row=row_idx, column=3, value='工作表中缺少必要的列（图号、部件名称、工序名称）')
        
        # 复制每个错误工作表的原始内容（重新以data_only模式加载以获取公式计算结果）
        wb_source_data_only = load_workbook('合金钢辙叉计件工价标准.xlsx', data_only=True)
        
        for error_info in error_sheets:
            sheet_name = error_info['sheet_name']
            # 使用data_only模式的工作表以获取公式计算结果
            original_ws = wb_source_data_only[sheet_name]
            
            # 创建新工作表（限制工作表名称长度）
            safe_name = sheet_name[:31] if len(sheet_name) > 31 else sheet_name
            new_ws = wb_error.create_sheet(safe_name)
            
            # 复制所有数据，确保公式被计算为结果值
            for row in original_ws.iter_rows():
                row_data = []
                for cell in row:
                    value = get_cell_value(cell)
                    row_data.append(value)
                new_ws.append(row_data)
            
            print(f"  已复制工作表: {sheet_name}")
        
        # 保存错误工作表文件
        error_file = 'error_sheets_export.xlsx'
        wb_error.save(error_file)
        
        print(f"错误工作表已导出到: {error_file}")
        print(f"包含 {len(error_sheets)} 个错误工作表和1个错误汇总表")
        
    except Exception as e:
        print(f"导出错误工作表时出错: {e}")


def write_to_template(data):
    """将数据写入模板文件"""
    print("正在写入模板文件...")
    
    try:
        # 加载模板文件
        wb_template = load_workbook('process_price_template.xlsx')
        
        # 使用模板中的"工序价格"工作表
        if '工序价格' in wb_template.sheetnames:
            ws = wb_template['工序价格']
        else:
            ws = wb_template.active
        
        # 清空现有数据（保留表头）
        max_row = ws.max_row
        if max_row > 1:
            ws.delete_rows(2, max_row - 1)
        
        # 模板列结构：工序编号、工序名称、部件、图号、型号、单价、生效日期、备注、价格类型、包含工序
        # 获取当前日期
        today = date.today()
        today_str = today.strftime('%Y-%m-%d')
        
        # 写入数据
        for row_idx, item in enumerate(data, 2):
            ws.cell(row=row_idx, column=1, value=item['process_code'])      # 工序编号*
            ws.cell(row=row_idx, column=2, value=item['process_name'])      # 工序名称*
            ws.cell(row=row_idx, column=3, value=item['part_name'])         # 部件
            ws.cell(row=row_idx, column=4, value=item['drawing_num'])       # 图号
            ws.cell(row=row_idx, column=5, value=item['model'])             # 型号
            ws.cell(row=row_idx, column=6, value=item['price'])             # 单价*
            ws.cell(row=row_idx, column=7, value=today_str)                 # 生效日期*（当前日期）
            ws.cell(row=row_idx, column=8, value=item['remark'])            # 备注
            ws.cell(row=row_idx, column=9, value=item['price_type'])        # 价格类型*
            ws.cell(row=row_idx, column=10, value=','.join(item['included_processes']))  # 包含工序
        
        # 保存文件
        import time
        timestamp = int(time.time())
        output_file = f'process_price_output_{timestamp}.xlsx'
        wb_template.save(output_file)
        
        print(f"数据已成功写入 {output_file}")
        print(f"共处理 {len(data)} 条记录")
        
        return True
        
    except Exception as e:
        print(f"写入模板文件时出错: {e}")
        return False


def main():
    """主函数"""
    print("=== 工价标准数据转换工具 ===")
    print()
    
    # 检查必要的依赖
    try:
        import pypinyin
    except ImportError:
        print("错误: 缺少 pypinyin 库，请运行: pip install pypinyin")
        return
    
    # 处理源数据
    data = process_source_data()
    
    if not data:
        print("未找到有效数据，程序结束")
        return
    
    print(f"\n总共处理了 {len(data)} 条记录")
    
    # 写入模板
    success = write_to_template(data)
    
    if success:
        print("\n转换完成！")
    else:
        print("\n转换失败！")


if __name__ == "__main__":
    main()
