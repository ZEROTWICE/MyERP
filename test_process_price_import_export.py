import os
import requests
import time
from datetime import datetime
from openpyxl import load_workbook, Workbook

# 应用程序URL
BASE_URL = 'http://localhost:5000'
LOGIN_URL = f'{BASE_URL}/auth/login'
TEMPLATE_URL = f'{BASE_URL}/process_prices/template'
IMPORT_URL = f'{BASE_URL}/process_prices/import'
EXPORT_URL = f'{BASE_URL}/process_prices/export'

# 登录凭据
USERNAME = 'admin'
PASSWORD = 'admin123'

# 测试数据目录
TEST_DIR = 'test_data'
TEMPLATE_PATH = os.path.join(TEST_DIR, 'process_price_template.xlsx')
TEST_DATA_PATH = os.path.join(TEST_DIR, 'process_price_test_data.xlsx')
EXPORT_PATH = os.path.join(TEST_DIR, 'process_price_export.xlsx')

def login_session():
    """创建一个已登录的会话"""
    session = requests.Session()
    # 获取CSRF令牌
    response = session.get(LOGIN_URL)
    # 从响应中提取CSRF令牌
    csrf_token = None
    if 'name="csrf_token" value="' in response.text:
        csrf_token = response.text.split('name="csrf_token" value="')[1].split('"')[0]
    
    # 登录
    if csrf_token:
        login_data = {
            'username': USERNAME,
            'password': PASSWORD,
            'csrf_token': csrf_token
        }
        response = session.post(LOGIN_URL, data=login_data, allow_redirects=True)
        if response.url != LOGIN_URL:
            print("登录成功")
            return session
    
    print("登录失败")
    return None

def download_template(session):
    """下载工序价格模板"""
    response = session.get(TEMPLATE_URL)
    if response.status_code == 200:
        with open(TEMPLATE_PATH, 'wb') as f:
            f.write(response.content)
        print(f"模板已下载到 {TEMPLATE_PATH}")
        return True
    else:
        print(f"下载模板失败: {response.status_code}")
        return False

def check_template():
    """检查下载的模板是否包含小计相关列"""
    if not os.path.exists(TEMPLATE_PATH):
        print("模板文件不存在")
        return False
    
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    
    # 检查表头是否包含'价格类型'和'包含工序'
    headers = [cell.value for cell in ws[1]]
    price_type_index = None
    included_processes_index = None
    
    for i, header in enumerate(headers):
        if header == '价格类型*':
            price_type_index = i
        elif header == '包含工序':
            included_processes_index = i
    
    if price_type_index is None or included_processes_index is None:
        print("模板中缺少'价格类型'或'包含工序'列")
        return False
    
    # 检查示例数据
    example_row = [cell.value for cell in ws[3]]  # 第三行应该是小计示例
    if price_type_index < len(example_row) and example_row[price_type_index] == 'subtotal':
        print("模板中包含小计示例")
    else:
        print("模板中未找到小计示例")
        return False
    
    print("模板检查通过，包含价格类型和包含工序列")
    return True

def prepare_test_data():
    """准备测试数据"""
    wb = Workbook()
    ws = wb.active
    ws.title = "工序价格"
    
    # 设置表头
    headers = ['工序编号*', '工序名称*', '部件', '图号', '型号', '单价*', '生效日期*', '备注', '价格类型*', '包含工序']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    # 添加普通工序
    ws.cell(row=2, column=1, value='TEST001')
    ws.cell(row=2, column=2, value='测试工序1')
    ws.cell(row=2, column=3, value='测试部件')
    ws.cell(row=2, column=4, value='图号A')
    ws.cell(row=2, column=5, value='型号A')
    ws.cell(row=2, column=6, value=10.0)
    ws.cell(row=2, column=7, value=datetime.now().strftime('%Y-%m-%d'))
    ws.cell(row=2, column=8, value='测试备注')
    ws.cell(row=2, column=9, value='normal')
    
    # 添加普通工序2
    ws.cell(row=3, column=1, value='TEST002')
    ws.cell(row=3, column=2, value='测试工序2')
    ws.cell(row=3, column=3, value='测试部件')
    ws.cell(row=3, column=4, value='图号B')
    ws.cell(row=3, column=5, value='型号B')
    ws.cell(row=3, column=6, value=15.0)
    ws.cell(row=3, column=7, value=datetime.now().strftime('%Y-%m-%d'))
    ws.cell(row=3, column=8, value='测试备注')
    ws.cell(row=3, column=9, value='normal')
    
    # 添加小计工序
    ws.cell(row=4, column=1, value='TEST003')
    ws.cell(row=4, column=2, value='测试小计')
    ws.cell(row=4, column=3, value='测试部件')
    ws.cell(row=4, column=4, value='图号C')
    ws.cell(row=4, column=5, value='型号C')
    ws.cell(row=4, column=6, value='')  # 小计会自动计算
    ws.cell(row=4, column=7, value=datetime.now().strftime('%Y-%m-%d'))
    ws.cell(row=4, column=8, value='小计测试')
    ws.cell(row=4, column=9, value='subtotal')
    ws.cell(row=4, column=10, value='TEST001,TEST002')  # 包含TEST001和TEST002
    
    wb.save(TEST_DATA_PATH)
    print(f"测试数据已保存到 {TEST_DATA_PATH}")
    return True

def import_test_data(session):
    """导入测试数据"""
    if not os.path.exists(TEST_DATA_PATH):
        print("测试数据文件不存在")
        return False
    
    # 构建multipart表单数据
    files = {'file': open(TEST_DATA_PATH, 'rb')}
    
    # 发送导入请求
    response = session.post(IMPORT_URL, files=files)
    
    # 关闭文件
    files['file'].close()
    
    if response.status_code == 200:
        try:
            result = response.json()
            if result.get('success'):
                print(f"导入成功: {result.get('message')}")
                return True
            else:
                print(f"导入失败: {result.get('message')}")
                return False
        except:
            print(f"解析响应失败: {response.text}")
            return False
    else:
        print(f"导入请求失败: {response.status_code}")
        return False

def export_process_prices(session):
    """导出工序价格数据"""
    # 准备表单数据
    data = {
        'submit': '导出'
    }
    
    # 发送导出请求
    response = session.post(EXPORT_URL, data=data)
    
    if response.status_code == 200:
        with open(EXPORT_PATH, 'wb') as f:
            f.write(response.content)
        print(f"数据已导出到 {EXPORT_PATH}")
        return True
    else:
        print(f"导出失败: {response.status_code}")
        return False

def check_export_data():
    """检查导出的数据是否包含小计相关信息"""
    if not os.path.exists(EXPORT_PATH):
        print("导出文件不存在")
        return False
    
    wb = load_workbook(EXPORT_PATH)
    ws = wb.active
    
    # 检查表头是否包含'价格类型'和'包含工序'
    headers = [cell.value for cell in ws[1]]
    price_type_index = None
    included_processes_index = None
    
    for i, header in enumerate(headers):
        if header == '价格类型':
            price_type_index = i
        elif header == '包含工序':
            included_processes_index = i
    
    if price_type_index is None or included_processes_index is None:
        print("导出数据中缺少'价格类型'或'包含工序'列")
        return False
    
    # 查找刚刚导入的测试小计数据
    subtotal_found = False
    included_processes_found = False
    
    for row in ws.iter_rows(min_row=2):  # 从第二行开始
        process_code = row[0].value
        price_type = row[price_type_index].value if price_type_index < len(row) else None
        included_processes = row[included_processes_index].value if included_processes_index < len(row) else None
        
        if process_code == 'TEST003' and price_type == 'subtotal':
            subtotal_found = True
            if included_processes and 'TEST001' in included_processes and 'TEST002' in included_processes:
                included_processes_found = True
    
    if subtotal_found and included_processes_found:
        print("导出数据包含小计信息和包含工序信息")
        return True
    elif subtotal_found:
        print("导出数据包含小计信息，但未找到包含工序信息")
        return False
    else:
        print("导出数据中未找到小计信息")
        return False

def run_tests():
    """运行所有测试"""
    # 确保测试目录存在
    os.makedirs(TEST_DIR, exist_ok=True)
    
    # 登录
    session = login_session()
    if not session:
        return
    
    # 测试1: 下载模板
    if not download_template(session):
        return
    
    # 测试2: 检查模板
    if not check_template():
        return
    
    # 测试3: 准备测试数据
    if not prepare_test_data():
        return
    
    # 测试4: 导入测试数据
    if not import_test_data(session):
        return
    
    # 等待数据处理完成
    print("等待数据处理...")
    time.sleep(2)
    
    # 测试5: 导出工序价格
    if not export_process_prices(session):
        return
    
    # 测试6: 检查导出数据
    if not check_export_data():
        return
    
    print("所有测试通过！小计功能在Excel导入导出中正常工作。")

if __name__ == '__main__':
    run_tests() 