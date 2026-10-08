# -*- coding: utf-8 -*-
"""合金钢辙叉 48 页表格（BOM）转换工具。

数据源：``工具/48页表格/`` 下的 48 个「（N）XXX.xlsx」（另有 ``1-24页表格.xlsx``
打印合订本与 Excel 锁文件 ``~$``，均不作为数据源）。

版式（逐文件核对后归纳）
------------------------
* 文件第 1 行标题为 ``总图号：XXX`` / ``总图号:XXX`` / ``图号：XXX``（全角与半角
  冒号混用，末尾可能带空格）。这是该文件所有 sheet 的父件图号。
* 表头行位置不一：多数在第 2 行，``（9）YHHC CZ577-03.xlsx`` 的 ``组合心轨明细表``
  在第 1 行（没有标题行）。
* 一个 sheet 可能**左右并排两张表**（``（18）（19）（21）（23）（24）``）：各自有独立
  的 ``序号`` 列，因此按 ``序号`` 列切成若干列组，每组独立解析后合并进同一张 BOM。
* 列组宽度（裁剪掉纯空列后）与版式的对应关系：

  ====  ====================================================  ====
  宽度  字段序                                                说明
  ====  ====================================================  ====
  8     ``序号|代码|图号|名称|数量|材料|单重|备注``            A 型，无单位、无总重
  8     ``序号|图号|名称|数量|材料|单重|总重|备注``            D 型，无单位、无代码
  9     ``序号|图号|名称|单位|数量|材料|单重|总重|备注``       B/E 型
  11    ``序号|代码|图号|名称|单位|数量|材料|单重|总重|备注|类型``  C 型
  ====  ====================================================  ====

* **表头单元格常有损坏**，因此列位以「版式序位」为准，表头文字只用于挑选版式与
  报告异常：``（11）YHHC CZ577-I-03.xlsx`` 的表头是 ``序号|图 号|名|鞋|数量|材料|
  kg)| |备注``（「名称」被拆成「名」「鞋」），``（12）YHCCZ2209-IV-03.xlsx`` 的数量
  列表头整个丢失，``（13）YHHC CZ2209-I-03.xlsx`` 的数量列表头只剩一个「量」。
* 数据质量问题：数量列用小写字母 ``l`` 冒充 1、数量为空（对称件/加宽件）、页脚文本
  混入数据区。这些一律**行级记录告警**，不中断整份文件。

输出
----
``convert_bom_folder(folder_path)`` / ``convert_bom_file(file_path)`` 返回 dict：

* ``data``：明细行列表，每行含 ``file_name`` / ``sheet_name`` / ``table_index`` /
  ``parent_drawing`` / ``parent_model`` / ``seq`` / ``code`` / ``drawing_no`` /
  ``name`` / ``unit`` / ``quantity`` / ``material`` / ``unit_weight`` /
  ``total_weight`` / ``notes`` / ``part_type`` / ``source_row``。
* ``parents``：父件列表，含 ``parent_drawing`` / ``parent_model`` / ``file_name`` /
  ``sheet_names`` / ``child_count``。
* ``error_sheets`` / ``warnings`` / 各计数键。
"""

import glob
import os
import re

from openpyxl import load_workbook

#: 允许的 Excel 扩展名
BOM_FILE_SUFFIXES = ('.xlsx', '.xlsm', '.xltx', '.xltm')
#: 不作为数据源的文件名前缀：Excel 锁文件、打印合订本、隐藏文件
BOM_SKIP_FILE_PREFIXES = ('~$', '1-24', '.')
#: 扫描列的上限，避免被空格撑大的 max_column 拖慢
BOM_MAX_COLUMNS = 24
#: 单个列组的最大宽度
BOM_MAX_GROUP_WIDTH = 11
#: 表头行只在这些行里找
BOM_HEADER_SCAN_ROWS = 8

ROMAN_MAP = {
    'Ⅰ': 'I', 'Ⅱ': 'II', 'Ⅲ': 'III', 'Ⅳ': 'IV', 'Ⅴ': 'V', 'Ⅵ': 'VI',
    'Ⅶ': 'VII', 'Ⅷ': 'VIII', 'Ⅸ': 'IX', 'Ⅹ': 'X',
    'ⅰ': 'i', 'ⅱ': 'ii', 'ⅲ': 'iii', 'ⅳ': 'iv', 'ⅴ': 'v',
    'Ｉ': 'I', 'Ｖ': 'V', 'Ｘ': 'X',
}

#: 字段别名。比较前统一做 `_norm`（去空白/全角空格/标点并 strip）
FIELD_ALIASES = {
    '序号': ('序号', '序'),
    '代码': ('代码', '物料编码', '编码'),
    '图号': ('图号', '图纸号', '图样号'),
    '名称': ('名称', '名称及规格', '品名', '名规格', '品名规格'),
    '单位': ('单位', '计量单位'),
    '数量': ('数量', '用量', '数', '量'),
    '材料': ('材料', '材质', '材料牌号'),
    '单重': ('单重', '单件重量', '单质量', '单件质量', '单体', '単质量', '鸿质量'),
    '总重': ('总重', '总质量', '总重量', '总里', '总体'),
    '备注': ('备注', '说明', '附注'),
    '类型': ('类型', '类别'),
}

#: 版式定义：宽度 -> ((字段序...), 说明)。同宽度多种版式时按表头文字打分挑选。
BOM_LAYOUTS = (
    (8, ('序号', '代码', '图号', '名称', '数量', '材料', '单重', '备注'), 'A型（无单位、无总重）'),
    (8, ('序号', '图号', '名称', '数量', '材料', '单重', '总重', '备注'), 'D型（无单位、无代码）'),
    (9, ('序号', '图号', '名称', '单位', '数量', '材料', '单重', '总重', '备注'), 'B/E型'),
    (11, ('序号', '代码', '图号', '名称', '单位', '数量', '材料', '单重', '总重', '备注', '类型'), 'C型'),
)

#: 页脚行关键字：命中则整行跳过
FOOTER_KEYWORDS = ('编制', '审核', '批准', '制表', '整理', '校对', '日期', '检验')
#: 只有「名称」列有值（图号为空）时的页脚关键字，避免把说明文字当成子件
FOOTER_NAME_KEYWORDS = ('总重', '总质量', '总体积', '说明', '备注', '注：', '注:')

#: 数量列里冒充 1 的字符
QUANTITY_ONE_CHARS = ('l', 'L', 'I', 'i', '|', '丨')
#: 无法解析的数量文本
QUANTITY_BAD_TEXTS = ('一', '二', '三', '四', '五', '六', '七', '八', '九', '十', '若干', '批')

_TITLE_RE = re.compile(r'图\s*号\s*[：:]\s*(?P<drawing>[^，,;；]*)')
_FILE_INDEX_RE = re.compile(r'^[（(]\s*\d+\s*[)）]\s*')


def _text(value):
    """把单元格值转成去空白字符串。"""
    if value is None:
        return ''
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _norm(value):
    """归一化：去掉空白（含全角空格）与常见标点，用于栏目名比较。"""
    text = _text(value)
    text = re.sub(r'[\s\u3000]+', '', text)
    text = re.sub(r'[:：()（）\[\]{}、,，.。/／\\\-_]+', '', text)
    return text.upper()


def _to_float(value):
    """宽松解析数值；无法解析返回 None。"""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value)
    if not text:
        return None
    match = re.search(r'-?\d+(?:\.\d+)?', text.replace(',', ''))
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def _field_of(value):
    """把表头单元格文字映射回规范字段名；识别不出返回 None。"""
    text = _norm(value)
    if not text:
        return None
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if _norm(alias) == text:
                return field
    best = None
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            alias_norm = _norm(alias)
            if len(alias_norm) >= 2 and alias_norm in text:
                if best is None or len(alias_norm) > len(_norm(best[1])):
                    best = (field, alias)
    return best[0] if best else None


class AlloySteelBomConverter(object):
    """合金钢辙叉 48 页表格转换器。"""

    # ---------------------------------------------------------------- 入口

    @staticmethod
    def convert_bom_folder(folder_path):
        """转换整个「48页表格」目录。

        :param folder_path: 目录路径
        :return: 结果 dict（见模块 docstring）
        """
        if not folder_path or not os.path.isdir(folder_path):
            return AlloySteelBomConverter._empty_result(
                success=False, message='目录不存在：{0}'.format(folder_path))

        file_paths = []
        for path in sorted(glob.glob(os.path.join(folder_path, '*'))):
            base = os.path.basename(path)
            if not os.path.isfile(path):
                continue
            if base.startswith(BOM_SKIP_FILE_PREFIXES):
                continue
            if os.path.splitext(base)[1].lower() not in BOM_FILE_SUFFIXES:
                continue
            file_paths.append(path)

        if not file_paths:
            return AlloySteelBomConverter._empty_result(
                success=False, message='目录中没有可解析的 Excel 文件（.xlsx）')

        return AlloySteelBomConverter._convert_files(file_paths)

    @staticmethod
    def convert_bom_file(file_path):
        """转换单个 BOM 文件。"""
        if not file_path or not os.path.isfile(file_path):
            return AlloySteelBomConverter._empty_result(
                success=False, message='文件不存在：{0}'.format(file_path))
        return AlloySteelBomConverter._convert_files([file_path])

    @staticmethod
    def convert_bom_files(file_paths):
        """转换一批 BOM 文件（HTTP 上传场景：文件先落到 uploads/temp 再传路径进来）。

        与 convert_bom_folder 的差别只在于「文件清单由调用方给」：上传文件名可能是
        `xxx.xlsx` 而没有「（N）」前缀，所以这里**不按「（N）」前缀过滤**。
        但**同一条文件级跳过表 `BOM_SKIP_FILE_PREFIXES` 仍然生效** —— 实测（2026-10-09）
        把目录里全部 49 份一次上传时，`1-24页表格.xlsx`（左右并排的打印合订本）会被
        当成一个真正的父件建出来，并带 9 条「未找到表头行／版式不识别」的 error_sheets；
        而它在 `convert_bom_folder` 里本来就被跳过。两条入口必须同口径。
        """
        if not file_paths:
            return AlloySteelBomConverter._empty_result(
                success=False, message='没有上传文件')

        selected = []
        for path in file_paths:
            if not path or not os.path.isfile(path):
                continue
            base = os.path.basename(path)
            if base.startswith(BOM_SKIP_FILE_PREFIXES):
                continue
            if os.path.splitext(base)[1].lower() not in BOM_FILE_SUFFIXES:
                continue
            selected.append(path)

        if not selected:
            return AlloySteelBomConverter._empty_result(
                success=False, message='没有可解析的 Excel 文件（.xlsx/.xlsm）')

        return AlloySteelBomConverter._convert_files(selected)

    # ------------------------------------------------------------ 内部实现

    @staticmethod
    def _empty_result(success, message, **extra):
        result = {
            'success': success,
            'message': message,
            'data': [],
            'parents': [],
            'error_sheets': [],
            'warnings': [],
            'file_count': 0,
            'sheet_count': 0,
            'total_count': 0,
            'error_count': 0,
        }
        result.update(extra)
        return result

    @staticmethod
    def _convert_files(file_paths):
        data = []
        parents = {}
        error_sheets = []
        warnings = []
        sheet_count = 0

        for path in file_paths:
            base = os.path.basename(path)
            file_result = AlloySteelBomConverter._parse_workbook(path, warnings)
            file_parent = file_result['parent_drawing']
            parent_model = file_result['parent_model']

            if not file_result['sheets']:
                error_sheets.append({
                    'file_name': base,
                    'sheet_name': '',
                    'error_type': '文件无法解析',
                    'detail': file_result['error'] or '没有可解析的工作表',
                })
                continue

            file_records = 0
            for sheet in file_result['sheets']:
                sheet_count += 1
                if sheet['error']:
                    error_sheets.append({
                        'file_name': base,
                        'sheet_name': sheet['sheet_name'],
                        'error_type': sheet['error'],
                        'detail': sheet['detail'],
                    })
                    continue
                parent_drawing = sheet['parent_drawing'] or file_parent or parent_model
                for item in sheet['items']:
                    item['file_name'] = base
                    item['sheet_name'] = sheet['sheet_name']
                    item['parent_drawing'] = parent_drawing
                    item['parent_model'] = parent_model
                    data.append(item)
                    file_records += 1

                parent = parents.setdefault(parent_drawing, {
                    'parent_drawing': parent_drawing,
                    'parent_model': parent_model,
                    'file_name': base,
                    'sheet_names': [],
                    'child_count': 0,
                })
                if sheet['sheet_name'] not in parent['sheet_names']:
                    parent['sheet_names'].append(sheet['sheet_name'])
                parent['child_count'] += len(sheet['items'])

            if not file_records and not file_result['parent_drawing']:
                error_sheets.append({
                    'file_name': base,
                    'sheet_name': '',
                    'error_type': '没有可导入的明细行',
                    'detail': '未找到标题行或表头行',
                })

        success = bool(data)
        message = '' if success else '没有解析到任何 BOM 明细行'
        result = {
            'success': success,
            'message': message,
            'data': data,
            'parents': [parents[key] for key in sorted(parents)],
            'error_sheets': error_sheets,
            'warnings': warnings,
            'file_count': len(file_paths),
            'sheet_count': sheet_count,
            'total_count': len(data),
            'error_count': len(error_sheets),
        }
        return result

    @staticmethod
    def _parse_workbook(path, warnings):
        """解析一个工作簿，返回 {'parent_drawing','parent_model','sheets','error'}。"""
        base = os.path.basename(path)
        parent_model = _FILE_INDEX_RE.sub('', os.path.splitext(base)[0]).strip()
        result = {
            'parent_drawing': '',
            'parent_model': parent_model,
            'sheets': [],
            'error': '',
        }
        try:
            workbook = load_workbook(path, data_only=True, read_only=False)
        except Exception as first_error:
            try:
                workbook = load_workbook(path, data_only=False, read_only=False)
            except Exception as second_error:
                result['error'] = '文件格式不受支持或文件已损坏：{0}（{1}）'.format(
                    first_error, second_error)
                return result

        try:
            for sheet_name in workbook.sheetnames:
                sheet = AlloySteelBomConverter._parse_sheet(
                    workbook[sheet_name], sheet_name, base, warnings)
                if sheet['parent_drawing'] and not result['parent_drawing']:
                    # 一个工作簿里只有第一张有标题的表定义父件图号
                    result['parent_drawing'] = sheet['parent_drawing']
                result['sheets'].append(sheet)
        finally:
            close = getattr(workbook, 'close', None)
            if close:
                try:
                    close()
                except Exception:
                    pass
        return result

    @staticmethod
    def _parse_sheet(worksheet, sheet_name, file_name, warnings):
        """解析一张工作表。"""
        sheet = {
            'sheet_name': sheet_name,
            'parent_drawing': '',
            'header_row': 0,
            'layout': '',
            'items': [],
            'error': '',
            'detail': '',
        }
        max_row = worksheet.max_row or 0
        max_col = min(worksheet.max_column or 0, BOM_MAX_COLUMNS)
        if max_row < 1 or max_col < 1:
            sheet['error'] = '空工作表'
            sheet['detail'] = '工作表没有任何单元格'
            return sheet

        rows = [[worksheet.cell(row=r, column=c).value for c in range(1, max_col + 1)]
                for r in range(1, min(max_row, BOM_HEADER_SCAN_ROWS) + 1)]

        title = ''
        for row_values in rows:
            blob = ' '.join(_text(v) for v in row_values if _text(v))
            if re.search(r'图\s*号\s*[：:]', blob):
                title = blob
                break
        if title:
            match = _TITLE_RE.search(title)
            drawing = match.group('drawing') if match else ''
            sheet['parent_drawing'] = re.sub(r'\s+', '', _text(drawing))

        header_row = AlloySteelBomConverter._find_header_row(rows)
        if not header_row:
            sheet['error'] = '未找到表头行'
            sheet['detail'] = '前 {0} 行里没有同时出现「序号」与「图号/名称」'.format(
                len(rows))
            return sheet
        sheet['header_row'] = header_row
        header_cells = [worksheet.cell(row=header_row, column=c).value
                        for c in range(1, max_col + 1)]

        groups = AlloySteelBomConverter._locate_groups(header_cells)
        if not groups:
            sheet['error'] = '未找到必要的列'
            sheet['detail'] = '表头行里没有「序号」列'
            return sheet

        layouts = []
        for seq_col, end_col in groups:
            layout = AlloySteelBomConverter._resolve_layout(
                worksheet, header_row, seq_col, end_col)
            if layout is None:
                sheet['error'] = '未识别的表头版式'
                sheet['detail'] = '第 {0} 列起的列组宽度 {1} 不在已知版式内'.format(
                    seq_col, end_col - seq_col + 1)
                return sheet
            layouts.append(layout)

        sheet['layout'] = ' + '.join(layout['name'] for layout in layouts)

        for index, layout in enumerate(layouts, 1):
            items = AlloySteelBomConverter._parse_group(
                worksheet, header_row, layout, max_row, max_col,
                sheet_name, index, warnings)
            for item in items:
                item['table_index'] = index
                item['parent_drawing'] = sheet['parent_drawing']
                sheet['items'].append(item)

        if not sheet['items']:
            sheet['error'] = '没有可导入的明细行'
            sheet['detail'] = '表头行 {0} 之后没有识别到任何子件'.format(header_row)
        return sheet

    @staticmethod
    def _find_header_row(rows):
        """在前若干行里找表头行：必须同时出现「序号」与（图号/名称/代码）的精确表头。"""
        seq_alias = _norm('序号')
        other_aliases = set()
        for key in ('图号', '名称', '代码'):
            for alias in FIELD_ALIASES[key]:
                other_aliases.add(_norm(alias))
        found = 0
        for index, row_values in enumerate(rows, 1):
            texts = {_norm(v) for v in row_values}
            if seq_alias in texts and (texts & other_aliases):
                found = index
        return found

    @staticmethod
    def _locate_groups(header_cells):
        """按「序号」列把表头切成若干列组，返回 [(起始列, 结束列)]（1-based）。

        结束列取下一个「序号」列的前一列；最后一组取表头行最后一个非空列。
        """
        seq_cols = [index + 1 for index, value in enumerate(header_cells)
                    if _norm(value) == _norm('序号')]
        if not seq_cols:
            return []
        groups = []
        for position, seq_col in enumerate(seq_cols):
            if position + 1 < len(seq_cols):
                end_col = seq_cols[position + 1] - 1
            else:
                non_empty = [index + 1 for index, value in enumerate(header_cells)
                             if _norm(value)]
                candidates = [c for c in non_empty if c >= seq_col]
                end_col = max(candidates or [seq_col])
            end_col = min(end_col, seq_col + BOM_MAX_GROUP_WIDTH - 1)
            if end_col < seq_col:
                continue
            groups.append((seq_col, end_col))
        return groups

    @staticmethod
    def _resolve_layout(worksheet, header_row, seq_col, end_col):
        """裁剪列组右侧的纯空列并按表头文字挑选版式。"""
        last_used = seq_col - 1
        for column in range(seq_col, end_col + 1):
            if AlloySteelBomConverter._column_has_value(worksheet, column, header_row):
                last_used = column
        width = last_used - seq_col + 1
        if width < 1:
            return None

        candidates = [entry for entry in BOM_LAYOUTS if entry[0] == width]
        if not candidates:
            return None

        scored = []
        for _, fields, name in candidates:
            score = 0
            for offset, field in enumerate(fields):
                cell = worksheet.cell(row=header_row, column=seq_col + offset).value
                text = _norm(cell)
                if not text:
                    continue
                aliases = tuple(_norm(a) for a in FIELD_ALIASES[field])
                if text in aliases:
                    score += 2
                elif any(len(a) >= 2 and a in text for a in aliases):
                    score += 1
                else:
                    other = _field_of(text)
                    if other and other != field:
                        score -= 1
            scored.append((score, name, fields))
        scored.sort(key=lambda entry: entry[0], reverse=True)
        best_score, best_name, best_fields = scored[0]
        return {
            'name': best_name,
            'fields': best_fields,
            'seq_col': seq_col,
            'end_col': last_used,
            'score': best_score,
            'ambiguous': len(scored) > 1 and scored[0][0] == scored[1][0],
        }

    @staticmethod
    def _column_has_value(worksheet, column, header_row):
        """列在表头行、子表头行或之后任一数据行有值即算「有值」。"""
        for row in range(header_row, (worksheet.max_row or header_row) + 1):
            if _text(worksheet.cell(row=row, column=column).value):
                return True
        return False

    @staticmethod
    def _parse_group(worksheet, header_row, layout, max_row, max_col,
                     sheet_name, table_index, warnings):
        """按版式逐行读取一个列组。"""
        seq_col = layout['seq_col']
        fields = layout['fields']
        col_of = {field: seq_col + offset for offset, field in enumerate(fields)}
        items = []
        ordinal = 0
        for row in range(header_row + 1, max_row + 1):
            seq_text = _text(worksheet.cell(row=row, column=seq_col).value)
            if _norm(seq_text) == _norm('序号'):
                # （23）右侧第 10 行重复了一次表头
                continue
            values = {}
            for field, column in col_of.items():
                if column > max_col:
                    values[field] = ''
                else:
                    values[field] = _text(worksheet.cell(row=row, column=column).value)

            name = values.get('名称', '')
            drawing_no = values.get('图号', '')
            code = values.get('代码', '')
            if not drawing_no and not name and not code:
                # 空行或子表头行（只有「单质量/总质量」）
                continue
            blob = ' '.join(v for v in values.values() if v)
            if any(keyword in blob for keyword in FOOTER_KEYWORDS):
                continue
            if not drawing_no and not code and any(
                    keyword in name for keyword in FOOTER_NAME_KEYWORDS):
                continue

            ordinal += 1
            quantity, quantity_warning = AlloySteelBomConverter._parse_quantity(
                values.get('数量', ''))
            if quantity_warning:
                warnings.append('{0} / {1} 第{2}行（{3}）{4}'.format(
                    sheet_name, layout['name'], row,
                    name or drawing_no or code, quantity_warning))

            items.append({
                'seq': seq_text or str(ordinal),
                'seq_text': seq_text,
                'code': code,
                'drawing_no': drawing_no,
                'name': name,
                'unit': values.get('单位', ''),
                'quantity': quantity,
                'quantity_text': values.get('数量', ''),
                'material': values.get('材料', ''),
                'unit_weight': _to_float(values.get('单重', '')),
                'total_weight': _to_float(values.get('总重', '')),
                'notes': values.get('备注', ''),
                'part_type': values.get('类型', ''),
                'source_row': row,
                'layout': layout['name'],
                'table_index': table_index,
                'file_name': '',
                'sheet_name': sheet_name,
                'parent_drawing': '',
                'parent_model': '',
            })
        return items

    @staticmethod
    def _parse_quantity(raw_text):
        """解析数量列，返回 (数量, 告警文案)。"""
        text = _text(raw_text)
        if not text:
            return None, '数量为空'
        if text in QUANTITY_ONE_CHARS:
            return 1.0, '数量「{0}」按 1 计'.format(text)
        if text in QUANTITY_BAD_TEXTS:
            return None, '数量「{0}」无法解析'.format(text)
        value = _to_float(text)
        if value is None:
            return None, '数量「{0}」无法解析'.format(text)
        return value, ''
