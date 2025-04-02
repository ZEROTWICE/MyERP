from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime
import os

class ExcelGenerator:
    @staticmethod
    def create_employee_template():
        """创建员工导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "员工信息"
        
        # 设置表头
        headers = ['工号', '姓名', '职位', '部门', '基本工资', '系数', '入职时间', '离职时间']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        # 添加日期格式说明
        ws.cell(row=2, column=7, value='格式：YYYY-MM-DD')
        ws.cell(row=2, column=8, value='格式：YYYY-MM-DD（可选）')
            
        return wb

    @staticmethod
    def create_process_price_template():
        """创建工序价格导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "工序价格"
        
        # 设置表头
        headers = ['工序编号', '工序名称', '部件', '图号', '型号', '单价', '生效日期', '备注']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        return wb

    @staticmethod
    def create_production_record_template():
        """创建生产记录导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "生产记录"
        
        # 设置表头
        headers = ['日期', '员工工号', '工序编号', '数量']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        return wb

    @staticmethod
    def create_bonus_penalty_template():
        """创建奖金/罚款导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "奖金和罚款记录"
        
        # 设置表头
        headers = ['日期', '员工工号', '类型(bonus/penalty)', '金额', '原因', '相关工序编号']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        return wb

    @staticmethod
    def create_task_template():
        """创建任务导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "任务分配"
        
        # 设置表头
        headers = ['员工工号', '工序编号', '数量', '目标完成日期', '备注']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        return wb

    @staticmethod
    def export_employees(employees):
        """导出员工数据"""
        wb = Workbook()
        ws = wb.active
        ws.title = "员工信息"
        
        # 设置表头
        headers = ['工号', '姓名', '职位', '部门', '基本工资', '系数', '入职时间', '离职时间', '状态']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 填充数据
        for row, employee in enumerate(employees, 2):
            ws.cell(row=row, column=1, value=employee.employee_id)
            ws.cell(row=row, column=2, value=employee.name)
            ws.cell(row=row, column=3, value=employee.position)
            ws.cell(row=row, column=4, value=employee.department)
            ws.cell(row=row, column=5, value=employee.base_salary)
            ws.cell(row=row, column=6, value=employee.coefficient)
            ws.cell(row=row, column=7, value=employee.hire_date.strftime('%Y-%m-%d') if employee.hire_date else '')
            ws.cell(row=row, column=8, value=employee.termination_date.strftime('%Y-%m-%d') if employee.termination_date else '')
            ws.cell(row=row, column=9, value=employee.status)
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        return wb

    @staticmethod
    def export_process_prices(process_prices):
        """导出工序价格数据"""
        wb = Workbook()
        ws = wb.active
        ws.title = "工序价格"
        
        # 设置表头
        headers = ['工序编号', '工序名称', '部件', '图号', '型号', '单价', '生效日期', '备注']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 填充数据
        for row, process in enumerate(process_prices, 2):
            ws.cell(row=row, column=1, value=process.process_code)
            ws.cell(row=row, column=2, value=process.process_name)
            ws.cell(row=row, column=3, value=process.component)
            ws.cell(row=row, column=4, value=process.drawing_no)
            ws.cell(row=row, column=5, value=process.model_no)
            ws.cell(row=row, column=6, value=process.price)
            ws.cell(row=row, column=7, value=process.effective_date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=8, value=process.notes)
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        return wb

    @staticmethod
    def export_production_records(records):
        """导出生产记录数据"""
        wb = Workbook()
        ws = wb.active
        ws.title = "生产记录"
        
        # 设置表头
        headers = ['日期', '员工工号', '员工姓名', '工序编号', '工序名称', '数量', '单价', '金额']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 填充数据
        for row, record in enumerate(records, 2):
            ws.cell(row=row, column=1, value=record.date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=2, value=record.employee.employee_id)
            ws.cell(row=row, column=3, value=record.employee.name)
            ws.cell(row=row, column=4, value=record.process.process_code)
            ws.cell(row=row, column=5, value=record.process.process_name)
            ws.cell(row=row, column=6, value=record.quantity)
            ws.cell(row=row, column=7, value=record.process.price)
            ws.cell(row=row, column=8, value=record.quantity * record.process.price)
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        return wb

    @staticmethod
    def export_bonus_penalties(records):
        """导出奖金/罚款数据"""
        wb = Workbook()
        ws = wb.active
        ws.title = "奖金和罚款记录"
        
        # 设置表头
        headers = ['日期', '员工工号', '员工姓名', '类型', '金额', '原因', '相关工序编号', '相关工序名称']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 填充数据
        for row, record in enumerate(records, 2):
            ws.cell(row=row, column=1, value=record.date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=2, value=record.employee.employee_id)
            ws.cell(row=row, column=3, value=record.employee.name)
            ws.cell(row=row, column=4, value='奖金' if record.type == 'bonus' else '罚款')
            ws.cell(row=row, column=5, value=float(record.amount))
            ws.cell(row=row, column=6, value=record.reason)
            if record.process:
                ws.cell(row=row, column=7, value=record.process.process_code)
                ws.cell(row=row, column=8, value=record.process.process_name)
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        return wb

    @staticmethod
    def export_tasks(tasks):
        """导出任务数据"""
        wb = Workbook()
        ws = wb.active
        ws.title = "任务分配"
        
        # 设置表头
        headers = ['分配日期', '员工工号', '员工姓名', '工序编号', '工序名称', '数量', '已完成数量', '完成率', '目标完成日期', '状态', '备注']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            
        # 填充数据
        for row, task in enumerate(tasks, 2):
            ws.cell(row=row, column=1, value=task.assigned_date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=2, value=task.employee.employee_id)
            ws.cell(row=row, column=3, value=task.employee.name)
            ws.cell(row=row, column=4, value=task.process.process_code)
            ws.cell(row=row, column=5, value=task.process.process_name)
            ws.cell(row=row, column=6, value=task.quantity)
            ws.cell(row=row, column=7, value=task.completed_quantity)
            ws.cell(row=row, column=8, value=f"{task.completion_rate:.1%}")
            ws.cell(row=row, column=9, value=task.target_date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=10, value=task.status)
            ws.cell(row=row, column=11, value=task.notes)
            
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        return wb

    @staticmethod
    def parse_employee_data(file_path):
        """解析员工导入数据"""
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):  # 跳过空行
                continue
                
            data.append({
                'employee_id': str(row[0]),
                'name': str(row[1]),
                'position': str(row[2]),
                'department': str(row[3]),
                'base_salary': float(row[4]),
                'coefficient': float(row[5])
            })
            
        return data

    @staticmethod
    def parse_process_price_data(file_path):
        """解析工序价格导入数据"""
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):  # 跳过空行
                continue
                
            data.append({
                'process_code': str(row[0]),
                'process_name': str(row[1]),
                'component': str(row[2]),
                'drawing_no': str(row[3]),
                'model_no': str(row[4]),
                'price': float(row[5]),
                'effective_date': datetime.strptime(str(row[6]), '%Y-%m-%d').date(),
                'notes': str(row[7]) if row[7] else None
            })
            
        return data

    @staticmethod
    def parse_production_record_data(file_path):
        """解析生产记录导入数据"""
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):  # 跳过空行
                continue
                
            data.append({
                'date': datetime.strptime(str(row[0]), '%Y-%m-%d').date(),
                'employee_id': str(row[1]),
                'process_code': str(row[2]),
                'quantity': int(row[3])
            })
            
        return data

    @staticmethod
    def parse_bonus_penalty_data(file_path):
        """解析奖金/罚款导入数据"""
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):  # 跳过空行
                continue
                
            data.append({
                'date': datetime.strptime(str(row[0]), '%Y-%m-%d').date(),
                'employee_id': str(row[1]),
                'type': str(row[2]).lower(),
                'amount': float(row[3]),
                'reason': str(row[4]),
                'process_code': str(row[5]) if row[5] else None
            })
            
        return data

    @staticmethod
    def parse_task_data(file_path):
        """解析任务导入数据"""
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):  # 跳过空行
                continue
                
            data.append({
                'employee_id': str(row[0]),
                'process_code': str(row[1]),
                'quantity': int(row[2]),
                'target_date': datetime.strptime(str(row[3]), '%Y-%m-%d').date(),
                'notes': str(row[4]) if row[4] else None
            })
            
        return data 