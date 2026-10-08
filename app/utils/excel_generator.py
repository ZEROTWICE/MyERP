from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

class ExcelGenerator:
    @staticmethod
    def _apply_header_style(cell):
        """应用表头样式"""
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    @staticmethod
    def _apply_cell_style(cell):
        """应用单元格样式"""
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    @staticmethod
    def create_employee_template():
        """创建员工导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "员工信息"
        
        # 设置表头
        headers = ['工号*', '姓名*', '职位*', '部门*', '基本工资*', '系数*', '入职时间*', '离职时间']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
            
        # 添加说明行
        ws.cell(row=2, column=1, value='EMP001')
        ws.cell(row=2, column=2, value='张三')
        ws.cell(row=2, column=3, value='技术员')
        ws.cell(row=2, column=4, value='生产部')
        ws.cell(row=2, column=5, value=5000)
        ws.cell(row=2, column=6, value=1.0)
        ws.cell(row=2, column=7, value='2024-01-01')
        ws.cell(row=2, column=8, value='')
        
        # 添加格式说明
        ws.cell(row=4, column=1, value='说明：')
        ws.cell(row=5, column=1, value='入职时间格式：YYYY-MM-DD（必填）')
        ws.cell(row=6, column=1, value='离职时间格式：YYYY-MM-DD（可选，留空表示在职）')
        ws.cell(row=7, column=1, value='带*号的字段为必填项')
        
        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        return wb

    @staticmethod
    def create_process_price_template():
        """创建工序价格导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "工序价格"
        
        # 设置表头
        headers = ['工序编号*', '工序名称*', '部件', '图号', '型号', '单价*', '生效日期*', '备注', '价格类型*', '包含工序']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
            
        # 添加说明行
        ws.cell(row=2, column=1, value='P001')
        ws.cell(row=2, column=2, value='车削')
        ws.cell(row=2, column=3, value='轴承')
        ws.cell(row=2, column=4, value='DWG001')
        ws.cell(row=2, column=5, value='M001')
        ws.cell(row=2, column=6, value=10.5)
        ws.cell(row=2, column=7, value='2024-01-01')
        ws.cell(row=2, column=8, value='标准工序')
        ws.cell(row=2, column=9, value='normal')
        ws.cell(row=2, column=10, value='')
        
        ws.cell(row=3, column=1, value='P002')
        ws.cell(row=3, column=2, value='小计工序')
        ws.cell(row=3, column=3, value='轴承')
        ws.cell(row=3, column=4, value='DWG002')
        ws.cell(row=3, column=5, value='M002')
        ws.cell(row=3, column=6, value='') # 小计会自动计算
        ws.cell(row=3, column=7, value='2024-01-01')
        ws.cell(row=3, column=8, value='小计工序')
        ws.cell(row=3, column=9, value='subtotal')
        ws.cell(row=3, column=10, value='P001,P003') # 指定要包含的工序编号，用逗号分隔
        
        # 添加说明
        ws.cell(row=5, column=1, value='说明')
        ws.cell(row=6, column=1, value='价格类型')
        ws.cell(row=6, column=2, value='normal - 普通工价；subtotal - 小计')
        ws.cell(row=7, column=1, value='包含工序')
        ws.cell(row=7, column=2, value='当价格类型为subtotal时，填写包含的工序编号，用逗号分隔。小计价格将自动计算为包含工序价格的总和。')
        
        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        return wb

    @staticmethod
    def create_production_record_template():
        """创建生产记录导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "生产记录"
        
        # 设置表头
        headers = ['日期*', '员工工号*', '工序编号*', '数量*']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
            
        # 添加说明行
        ws.cell(row=2, column=1, value='2024-01-01')
        ws.cell(row=2, column=2, value='EMP001')
        ws.cell(row=2, column=3, value='P001')
        ws.cell(row=2, column=4, value=100)
        
        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        return wb

    @staticmethod
    def create_bonus_penalty_template():
        """创建奖惩记录导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "奖惩记录"
        
        # 设置表头
        headers = ['日期*', '员工工号*', '类型*', '金额*', '工序编号', '原因']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
            
        # 添加说明行
        ws.cell(row=2, column=1, value='2024-01-01')
        ws.cell(row=2, column=2, value='EMP001')
        ws.cell(row=2, column=3, value='bonus')  # 或 penalty
        ws.cell(row=2, column=4, value=100)
        ws.cell(row=2, column=5, value='P001')  # 可选
        ws.cell(row=2, column=6, value='优秀表现奖励')
        
        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        return wb

    @staticmethod
    def create_task_template():
        """创建任务导入模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "任务分配"
        
        # 设置表头
        headers = ['员工工号*', '工序编号*', '数量*', '目标日期*', '备注']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
            
        # 添加说明行
        ws.cell(row=2, column=1, value='EMP001')
        ws.cell(row=2, column=2, value='P001')
        ws.cell(row=2, column=3, value=100)
        ws.cell(row=2, column=4, value='2024-01-31')
        ws.cell(row=2, column=5, value='紧急任务')
        
        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        return wb

    @staticmethod
    def parse_employee_data(file_path):
        """解析员工导入数据"""
        from openpyxl import load_workbook
        from datetime import datetime, date
        
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row_num, row in enumerate(ws.iter_rows(min_row=2), start=2):  # 跳过表头
            if not any(cell.value for cell in row):  # 跳过空行
                continue
                
            # 跳过说明行（从第4行开始的说明）
            if row_num >= 4 and row[0].value and str(row[0].value).startswith('说明'):
                continue
                
            try:
                # 处理入职时间
                hire_date_value = row[6].value if len(row) > 6 else None
                hire_date = None
                if hire_date_value:
                    if isinstance(hire_date_value, date):
                        hire_date = hire_date_value
                    elif isinstance(hire_date_value, str):
                        hire_date = datetime.strptime(hire_date_value.strip(), '%Y-%m-%d').date()
                
                # 处理离职时间
                termination_date_value = row[7].value if len(row) > 7 else None
                termination_date = None
                if termination_date_value:
                    if isinstance(termination_date_value, date):
                        termination_date = termination_date_value
                    elif isinstance(termination_date_value, str) and termination_date_value.strip():
                        termination_date = datetime.strptime(termination_date_value.strip(), '%Y-%m-%d').date()
                
                data.append({
                'employee_id': str(row[0].value).strip(),
                'name': str(row[1].value).strip(),
                'position': str(row[2].value).strip(),
                'department': str(row[3].value).strip(),
                'base_salary': float(row[4].value),
                    'coefficient': float(row[5].value),
                    'hire_date': hire_date,
                    'termination_date': termination_date
            })
            except (ValueError, IndexError, TypeError) as e:
                # 记录解析错误，但继续处理其他行
                print(f"第{row_num}行数据解析失败: {e}")
                continue
        
        return data

    @staticmethod
    def parse_process_price_data(file_path):
        """解析工序价格导入数据"""
        from openpyxl import load_workbook
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row in ws.iter_rows(min_row=2):  # 跳过表头
            if not any(cell.value for cell in row):  # 跳过空行
                continue
            
            # 检查行中是否有足够的单元格
            if len(row) < 9:
                continue
            
            price_type = 'normal'
            included_processes = []
            
            # 获取价格类型
            if len(row) >= 9 and row[8].value:
                price_type = str(row[8].value).strip().lower()
                if price_type not in ['normal', 'subtotal']:
                    price_type = 'normal'
            
            # 获取包含工序
            if price_type == 'subtotal' and len(row) >= 10 and row[9].value:
                included_processes = [code.strip() for code in str(row[9].value).split(',') if code.strip()]
            
            data.append({
                'process_code': str(row[0].value).strip(),
                'process_name': str(row[1].value).strip(),
                'component': str(row[2].value).strip() if row[2].value else None,
                'drawing_no': str(row[3].value).strip() if row[3].value else None,
                'model_no': str(row[4].value).strip() if row[4].value else None,
                'price': float(row[5].value) if row[5].value else 0,
                'effective_date': row[6].value,
                'notes': str(row[7].value).strip() if row[7].value else None,
                'price_type': price_type,
                'included_processes': included_processes
            })
        
        return data

    @staticmethod
    def parse_production_record_data(file_path):
        """解析生产记录导入数据"""
        from openpyxl import load_workbook
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row in ws.iter_rows(min_row=2):  # 跳过表头
            if not any(cell.value for cell in row):  # 跳过空行
                continue
            data.append({
                'date': row[0].value,
                'employee_id': str(row[1].value).strip(),
                'process_code': str(row[2].value).strip(),
                'quantity': int(row[3].value),
                # 模板只有 4 列（日期/工号/工序/数量），备注为可选第 5 列；缺失时取 None，
                # 保持与路由 notes=data['notes'] 的键约定一致（缺键会 KeyError 致整批导入失败）
                'notes': str(row[4].value).strip() if len(row) >= 5 and row[4].value else None
            })
        
        return data

    @staticmethod
    def parse_bonus_penalty_data(file_path):
        """解析奖惩记录导入数据"""
        from openpyxl import load_workbook
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row in ws.iter_rows(min_row=2):  # 跳过表头
            if not any(cell.value for cell in row):  # 跳过空行
                continue
            data.append({
                'date': row[0].value,
                'employee_id': str(row[1].value).strip(),
                'type': str(row[2].value).strip(),
                'amount': float(row[3].value),
                'process_code': str(row[4].value).strip() if row[4].value else None,
                'reason': str(row[5].value).strip() if row[5].value else None
            })
        
        return data

    @staticmethod
    def parse_task_data(file_path):
        """解析任务导入数据"""
        from openpyxl import load_workbook
        wb = load_workbook(file_path)
        ws = wb.active
        data = []
        
        for row in ws.iter_rows(min_row=2):  # 跳过表头
            if not any(cell.value for cell in row):  # 跳过空行
                continue
            data.append({
                'employee_id': str(row[0].value).strip(),
                'process_code': str(row[1].value).strip(),
                'quantity': int(row[2].value),
                'target_date': row[3].value,
                'notes': str(row[4].value).strip() if row[4].value else None
            })
        
        return data

    @staticmethod
    def export_employees(employees):
        """导出员工信息"""
        wb = Workbook()
        ws = wb.active
        ws.title = "员工信息"

        # 设置表头
        headers = ['工号', '姓名', '职位', '部门', '基本工资', '系数', '入职时间', '离职时间', '状态', '是否管理员']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)

        # 写入数据
        for row, employee in enumerate(employees, 2):
            ws.cell(row=row, column=1, value=employee.employee_id)
            ws.cell(row=row, column=2, value=employee.name)
            ws.cell(row=row, column=3, value=employee.position)
            ws.cell(row=row, column=4, value=employee.department)
            ws.cell(row=row, column=5, value=employee.base_salary)
            ws.cell(row=row, column=6, value=employee.coefficient)
            ws.cell(row=row, column=7, value=employee.hire_date.strftime('%Y-%m-%d') if employee.hire_date else '')
            ws.cell(row=row, column=8, value=employee.termination_date.strftime('%Y-%m-%d') if employee.termination_date else '')
            ws.cell(row=row, column=9, value=employee.status)
            ws.cell(row=row, column=10, value='是' if employee.user and employee.user.role == 'admin' else '否')

        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2

        return wb

    @staticmethod
    def export_process_prices(process_prices):
        """导出工序价格"""
        wb = Workbook()
        ws = wb.active
        ws.title = "工序价格"

        # 设置表头
        headers = ['工序编号', '工序名称', '部件', '图号', '型号', '单价', '生效日期', '备注', '价格类型', '包含工序']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)

        # 写入数据
        for row, process in enumerate(process_prices, 2):
            ws.cell(row=row, column=1, value=process.process_code)
            ws.cell(row=row, column=2, value=process.process_name)
            ws.cell(row=row, column=3, value=process.component)
            ws.cell(row=row, column=4, value=process.drawing_no)
            ws.cell(row=row, column=5, value=process.model_no)
            ws.cell(row=row, column=6, value=process.price)
            ws.cell(row=row, column=7, value=process.effective_date.strftime('%Y-%m-%d') if process.effective_date else '')
            ws.cell(row=row, column=8, value=process.notes)
            ws.cell(row=row, column=9, value=process.price_type)

            # 如果是小计，添加包含的工序
            if process.price_type == 'subtotal':
                included_process_codes = [group.process.process_code
                                          for group in process.included_processes
                                          if group.process]
                ws.cell(row=row, column=10, value=','.join(included_process_codes))

        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2

        return wb

    @staticmethod
    def export_production_records(records):
        """导出生产记录"""
        wb = Workbook()
        ws = wb.active
        ws.title = "生产记录"

        # 设置表头
        headers = ['日期', '员工工号', '员工姓名', '工序编号', '工序名称', '部件', '图号', '型号', '数量']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)

        # 写入数据
        for row, record in enumerate(records, 2):
            ws.cell(row=row, column=1, value=record.date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=2, value=record.employee.employee_id)
            ws.cell(row=row, column=3, value=record.employee.name)
            ws.cell(row=row, column=4, value=record.process.process_code)
            ws.cell(row=row, column=5, value=record.process.process_name)
            ws.cell(row=row, column=6, value=record.process.component)
            ws.cell(row=row, column=7, value=record.process.drawing_no)
            ws.cell(row=row, column=8, value=record.process.model_no)
            ws.cell(row=row, column=9, value=record.quantity)

        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2

        return wb

    @staticmethod
    def export_bonus_penalties(records):
        """导出奖惩记录"""
        wb = Workbook()
        ws = wb.active
        ws.title = "奖惩记录"

        # 设置表头
        headers = ['日期', '员工工号', '员工姓名', '类型', '金额', '工序编号', '原因']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)

        # 写入数据
        for row, record in enumerate(records, 2):
            ws.cell(row=row, column=1, value=record.date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=2, value=record.employee.employee_id)
            ws.cell(row=row, column=3, value=record.employee.name)
            ws.cell(row=row, column=4, value='奖金' if record.type == 'bonus' else '罚款')
            ws.cell(row=row, column=5, value=record.amount)
            ws.cell(row=row, column=6, value=record.process.process_code if record.process else '')
            ws.cell(row=row, column=7, value=record.reason)

        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2

        return wb

    @staticmethod
    def export_tasks(tasks):
        """导出任务记录"""
        wb = Workbook()
        ws = wb.active
        ws.title = "任务记录"

        # 设置表头
        headers = ['任务编号', '员工工号', '员工姓名', '工序编号', '工序名称', '数量', '目标日期', '状态', '分配日期', '完成日期', '备注']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)

        # 写入数据
        for row, task in enumerate(tasks, 2):
            ws.cell(row=row, column=1, value=task.id)
            ws.cell(row=row, column=2, value=task.employee.employee_id)
            ws.cell(row=row, column=3, value=task.employee.name)
            ws.cell(row=row, column=4, value=task.process.process_code)
            ws.cell(row=row, column=5, value=task.process.process_name)
            ws.cell(row=row, column=6, value=task.quantity)
            ws.cell(row=row, column=7, value=task.target_date.strftime('%Y-%m-%d') if task.target_date else '')
            ws.cell(row=row, column=8, value=task.status)
            ws.cell(row=row, column=9, value=task.assigned_date.strftime('%Y-%m-%d') if task.assigned_date else '')
            ws.cell(row=row, column=10, value=task.completed_at.strftime('%Y-%m-%d') if task.completed_at else '')
            ws.cell(row=row, column=11, value=task.notes)

        # 调整列宽
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2

        return wb

    @staticmethod
    def add_headers(wb, headers):
        """添加表头"""
        ws = wb.active
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
        return ws

    @staticmethod
    def add_row(ws, row_data, row_number=None):
        """添加数据行"""
        if row_number is None:
            row_number = ws.max_row + 1
        
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_number, column=col, value=value)
            ExcelGenerator._apply_cell_style(cell)

    @staticmethod
    def generate_response(wb, filename):
        """生成Excel响应"""
        from io import BytesIO
        from flask import send_file
        
        # 保存到内存
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        
        # 返回文件
        return send_file(
            excel_file,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        ) 

    @staticmethod
    def create_product_template():
        """创建产品（含BOM和工序）导入模板"""
        wb = Workbook()
        
        # 第一个工作表：产品信息
        ws_product = wb.active
        ws_product.title = "产品信息"
        
        # 产品信息表头
        product_headers = ['产品编码*', '产品名称*', '图号', '型号', '规格说明', '单位', '类别', '版本', '状态', '编码规则', '备注']
        for col, header in enumerate(product_headers, 1):
            cell = ws_product.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
            
        # 产品信息示例
        product_example = ['PROD001', '示例产品', 'DWG001', 'Model001', '示例规格说明', '件', '电子产品', '1.0', 'active', '产品编码规则', '示例产品备注']
        for col, value in enumerate(product_example, 1):
            ws_product.cell(row=2, column=col, value=value)
        
        # 产品信息说明
        ws_product.cell(row=4, column=1, value='说明')
        ws_product.cell(row=5, column=1, value='产品编码')
        ws_product.cell(row=5, column=2, value='必填，产品的唯一标识，必须与BOM表和工序表中的产品编码一致')
        ws_product.cell(row=6, column=1, value='状态')
        ws_product.cell(row=6, column=2, value='active-启用，inactive-停用，obsolete-淘汰')
        ws_product.cell(row=7, column=1, value='单位')
        ws_product.cell(row=7, column=2, value='默认为"件"')
        ws_product.cell(row=8, column=1, value='版本')
        ws_product.cell(row=8, column=2, value='默认为"1.0"')
        ws_product.cell(row=9, column=1, value='编码规则')
        ws_product.cell(row=9, column=2, value='填写编码规则名称，用于自动生成产品编码（可选）')
        
        # 第二个工作表：BOM物料清单
        ws_bom = wb.create_sheet(title="BOM物料清单")
        
        # BOM表头
        bom_headers = ['产品编码*', '物料类型*', '物料编码*', '物料名称', '用量*', '单位', '单价', '损耗率%', '备注', '序号']
        for col, header in enumerate(bom_headers, 1):
            cell = ws_bom.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
        
        # BOM示例数据
        bom_examples = [
            ['PROD001', 'raw', '内部编号001', '钢材', 2.5, 'kg', 15.0, 5.0, '主要原材料', 1],
            ['PROD001', 'finished', '成品编号001', '轴承', 1, '个', 50.0, 2.0, '成品配件', 2],
            ['PROD001', 'product', 'SUB001', '子产品', 1, '件', 100.0, 0, '子产品组件', 3]
        ]
        for row_idx, example in enumerate(bom_examples, 2):
            for col, value in enumerate(example, 1):
                ws_bom.cell(row=row_idx, column=col, value=value)
        
        # BOM说明
        ws_bom.cell(row=6, column=1, value='说明')
        ws_bom.cell(row=7, column=1, value='产品编码')
        ws_bom.cell(row=7, column=2, value='必填，必须与产品信息表中的产品编码一致')
        ws_bom.cell(row=8, column=1, value='物料类型')
        ws_bom.cell(row=8, column=2, value='raw-原材料，finished-成品，product-产品')
        ws_bom.cell(row=9, column=1, value='物料编码')
        ws_bom.cell(row=9, column=2, value='原材料使用内部编号，成品使用产品编号，产品使用产品编码')
        ws_bom.cell(row=10, column=1, value='损耗率')
        ws_bom.cell(row=10, column=2, value='填写百分比数值，如5表示5%损耗率')
        ws_bom.cell(row=11, column=1, value='序号')
        ws_bom.cell(row=11, column=2, value='BOM中的排序序号，数字越小排序越靠前')
        
        # 第三个工作表：产品工序
        ws_process = wb.create_sheet(title="产品工序")
        
        # 工序表头
        process_headers = ['产品编码*', '工序编号*', '工序名称', '序号*', '加工数量', '单价', '准备时间(分钟)', '加工时间(分钟)', '是否必需', '备注']
        for col, header in enumerate(process_headers, 1):
            cell = ws_process.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
        
        # 工序示例数据
        process_examples = [
            ['PROD001', 'P001', '车削', 1, 1, 15.0, 10, 30, True, '第一道工序'],
            ['PROD001', 'P002', '铣削', 2, 1, 20.0, 15, 45, True, '第二道工序'],
            ['PROD001', 'P003', '检验', 3, 1, 10.0, 5, 15, False, '质量检验']
        ]
        for row_idx, example in enumerate(process_examples, 2):
            for col, value in enumerate(example, 1):
                ws_process.cell(row=row_idx, column=col, value=value)
        
        # 工序说明
        ws_process.cell(row=6, column=1, value='说明')
        ws_process.cell(row=7, column=1, value='产品编码')
        ws_process.cell(row=7, column=2, value='必填，必须与产品信息表中的产品编码一致')
        ws_process.cell(row=8, column=1, value='工序编号')
        ws_process.cell(row=8, column=2, value='必填，必须是系统中已存在的工序编号')
        ws_process.cell(row=9, column=1, value='序号')
        ws_process.cell(row=9, column=2, value='工序执行顺序，数字越小执行越早')
        ws_process.cell(row=10, column=1, value='是否必需')
        ws_process.cell(row=10, column=2, value='TRUE-必需工序，FALSE-可选工序')
        ws_process.cell(row=11, column=1, value='单价')
        ws_process.cell(row=11, column=2, value='留空则使用工序价格表中的价格')
        
        # 调整所有工作表的列宽
        for ws in [ws_product, ws_bom, ws_process]:
            for col in ws.columns:
                max_length = 0
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        return wb

    @staticmethod
    def parse_product_data(file_path):
        """解析产品（含BOM和工序）导入数据"""
        from openpyxl import load_workbook
        from app.models import RawMaterial, FinishedProduct, ProcessPrice, CodeRule
        
        wb = load_workbook(file_path)
        result = {
            'products': [],
            'bom_items': [],
            'process_items': [],
            'errors': []
        }
        
        # 解析产品信息
        if '产品信息' in wb.sheetnames:
            ws_product = wb['产品信息']
            for row_num, row in enumerate(ws_product.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row) or (row[0] and str(row[0]).startswith('说明')):  # 跳过空行和说明行
                    continue
                    
                try:
                    # 处理编码规则
                    code_rule_name = str(row[9]).strip() if row[9] else None
                    code_rule_id = None
                    if code_rule_name:
                        # 查找编码规则
                        code_rule = CodeRule.query.filter_by(name=code_rule_name, code_type='product', is_active=True).first()
                        if code_rule:
                            code_rule_id = code_rule.id
                        else:
                            # 编码规则是可选的，记录警告但不跳过整行
                            result['errors'].append(f'产品信息第{row_num}行：找不到名为"{code_rule_name}"的产品编码规则，已忽略该编码规则')
                    
                    product_data = {
                        'product_code': str(row[0]).strip() if row[0] else None,
                        'product_name': str(row[1]).strip() if row[1] else None,
                        'drawing_number': str(row[2]).strip() if row[2] else None,
                        'model': str(row[3]).strip() if row[3] else None,
                        'specification': str(row[4]).strip() if row[4] else None,
                        'unit': str(row[5]).strip() if row[5] else '件',
                        'category': str(row[6]).strip() if row[6] else None,
                        'version': str(row[7]).strip() if row[7] else '1.0',
                        'status': str(row[8]).strip() if row[8] else 'active',
                        'code_rule_id': code_rule_id,
                        'notes': str(row[10]).strip() if row[10] else None
                    }
                    
                    # 验证必填字段
                    if not product_data['product_code']:
                        result['errors'].append(f'产品信息第{row_num}行：产品编码不能为空')
                        continue
                    if not product_data['product_name']:
                        result['errors'].append(f'产品信息第{row_num}行：产品名称不能为空')
                        continue
                    
                    result['products'].append(product_data)
                    
                except Exception as e:
                    result['errors'].append(f'产品信息第{row_num}行解析错误：{str(e)}')
        
        # 解析BOM物料清单
        if 'BOM物料清单' in wb.sheetnames:
            ws_bom = wb['BOM物料清单']
            for row_num, row in enumerate(ws_bom.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row) or (row[0] and str(row[0]).startswith('说明')):  # 跳过空行和说明行
                    continue
                    
                try:
                    # 获取物料信息
                    material_type = str(row[1]).strip().lower() if row[1] else None
                    material_code = str(row[2]).strip() if row[2] else None
                    material_name = str(row[3]).strip() if row[3] else None
                    
                    # 验证物料类型
                    if material_type not in ['raw', 'finished', 'product']:
                        result['errors'].append(f'BOM第{row_num}行：物料类型必须是raw、finished或product')
                        continue
                    
                    # 查找物料ID
                    material_id = None
                    if material_type == 'raw':
                        material = RawMaterial.query.filter_by(internal_number=material_code).first()
                        if material:
                            material_id = material.id
                        else:
                            result['errors'].append(f'BOM第{row_num}行：找不到内部编号为"{material_code}"的原材料')
                            continue
                    elif material_type == 'finished':
                        material = FinishedProduct.query.filter_by(product_number=material_code).first()
                        if material:
                            material_id = material.id
                        else:
                            result['errors'].append(f'BOM第{row_num}行：找不到产品编号为"{material_code}"的成品')
                            continue
                    elif material_type == 'product':
                        # 对于产品类型，我们先记录编码，稍后在导入时处理
                        material_id = material_code  # 临时使用编码作为ID
                    
                    bom_data = {
                        'product_code': str(row[0]).strip() if row[0] else None,
                        'material_type': material_type,
                        'material_id': material_id,
                        'material_code': material_code,  # 保留编码用于后续处理
                        'quantity': float(row[4]) if row[4] else 0,
                        'unit': str(row[5]).strip() if row[5] else '件',
                        'unit_cost': float(row[6]) if row[6] else 0,
                        'waste_rate': float(row[7]) if row[7] else 0,
                        'notes': str(row[8]).strip() if row[8] else None,
                        'sequence': int(row[9]) if row[9] else 0
                    }
                    
                    # 验证必填字段
                    if not bom_data['product_code']:
                        result['errors'].append(f'BOM第{row_num}行：产品编码不能为空')
                        continue
                    if not material_code:
                        result['errors'].append(f'BOM第{row_num}行：物料编码不能为空')
                        continue
                    if bom_data['quantity'] <= 0:
                        result['errors'].append(f'BOM第{row_num}行：用量必须大于0')
                        continue
                    
                    result['bom_items'].append(bom_data)
                    
                except Exception as e:
                    result['errors'].append(f'BOM第{row_num}行解析错误：{str(e)}')
        
        # 解析产品工序
        if '产品工序' in wb.sheetnames:
            ws_process = wb['产品工序']
            for row_num, row in enumerate(ws_process.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row) or (row[0] and str(row[0]).startswith('说明')):  # 跳过空行和说明行
                    continue
                    
                try:
                    process_code = str(row[1]).strip() if row[1] else None
                    
                    # 查找工序
                    process = ProcessPrice.query.filter_by(process_code=process_code, is_current=True).first()
                    if not process:
                        result['errors'].append(f'工序第{row_num}行：找不到工序编号为"{process_code}"的工序')
                        continue
                    
                    # 处理是否必需字段
                    is_required = True  # 默认值
                    if row[8] is not None:
                        is_required_str = str(row[8]).strip().lower()
                        if is_required_str in ['false', '否', 'no', '0']:
                            is_required = False
                        elif is_required_str in ['true', '是', 'yes', '1']:
                            is_required = True
                    
                    process_data = {
                        'product_code': str(row[0]).strip() if row[0] else None,
                        'process_id': process.id,
                        'process_code': process_code,
                        'sequence': int(row[3]) if row[3] else 0,
                        'quantity': int(row[4]) if row[4] else 1,
                        'unit_price': float(row[5]) if row[5] else None,
                        'setup_time': float(row[6]) if row[6] else 0,
                        'process_time': float(row[7]) if row[7] else 0,
                        'is_required': is_required,
                        'notes': str(row[9]).strip() if row[9] else None
                    }
                    
                    # 验证必填字段
                    if not process_data['product_code']:
                        result['errors'].append(f'工序第{row_num}行：产品编码不能为空')
                        continue
                    if process_data['sequence'] <= 0:
                        result['errors'].append(f'工序第{row_num}行：序号必须大于0')
                        continue
                    
                    result['process_items'].append(process_data)
                    
                except Exception as e:
                    result['errors'].append(f'工序第{row_num}行解析错误：{str(e)}')
        
        return result 

    @staticmethod
    def export_products(products):
        """导出产品信息（含BOM和工序）"""
        wb = Workbook()
        
        # 第一个工作表：产品信息
        ws_product = wb.active
        ws_product.title = "产品信息"

        # 产品信息表头
        product_headers = ['产品编码', '产品名称', '图号', '型号', '规格说明', '单位', '类别', '版本', '状态', '编码规则', '总成本', '创建时间', '备注']
        for col, header in enumerate(product_headers, 1):
            cell = ws_product.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)

        # 写入产品数据
        for row, product in enumerate(products, 2):
            ws_product.cell(row=row, column=1, value=product.product_code)
            ws_product.cell(row=row, column=2, value=product.product_name)
            ws_product.cell(row=row, column=3, value=product.drawing_number)
            ws_product.cell(row=row, column=4, value=product.model)
            ws_product.cell(row=row, column=5, value=product.specification)
            ws_product.cell(row=row, column=6, value=product.unit)
            ws_product.cell(row=row, column=7, value=product.category)
            ws_product.cell(row=row, column=8, value=product.version)
            ws_product.cell(row=row, column=9, value=product.status)
            # 获取编码规则名称
            code_rule_name = product.code_rule.name if product.code_rule else ''
            ws_product.cell(row=row, column=10, value=code_rule_name)
            ws_product.cell(row=row, column=11, value=product.total_cost)
            ws_product.cell(row=row, column=12, value=product.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws_product.cell(row=row, column=13, value=product.notes)

        # 第二个工作表：BOM物料清单
        ws_bom = wb.create_sheet(title="BOM物料清单")
        
        # BOM表头
        bom_headers = ['产品编码', '物料类型', '物料编码', '物料名称', '用量', '单位', '单价', '损耗率%', '实际用量', '总成本', '备注', '序号']
        for col, header in enumerate(bom_headers, 1):
            cell = ws_bom.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
        
        # 写入BOM数据
        bom_row = 2
        for product in products:
            for bom_item in product.bom_items.order_by('sequence'):
                ws_bom.cell(row=bom_row, column=1, value=product.product_code)
                ws_bom.cell(row=bom_row, column=2, value=bom_item.material_type)
                
                # 获取物料编码和名称
                material_code = ''
                material_name = ''
                if bom_item.material_type == 'raw':
                    material = bom_item.material_info
                    if material:
                        material_code = material.internal_number
                        material_name = material.material_name
                elif bom_item.material_type == 'finished':
                    material = bom_item.material_info
                    if material:
                        material_code = material.product_number
                        material_name = material.product_number
                elif bom_item.material_type == 'product':
                    material = bom_item.material_info
                    if material:
                        material_code = material.product_code
                        material_name = f"{material.product_name} ({material.product_code})"
                
                ws_bom.cell(row=bom_row, column=3, value=material_code)
                ws_bom.cell(row=bom_row, column=4, value=material_name)
                ws_bom.cell(row=bom_row, column=5, value=bom_item.quantity)
                ws_bom.cell(row=bom_row, column=6, value=bom_item.unit)
                ws_bom.cell(row=bom_row, column=7, value=bom_item.unit_cost)
                ws_bom.cell(row=bom_row, column=8, value=bom_item.waste_rate)
                ws_bom.cell(row=bom_row, column=9, value=bom_item.actual_quantity)
                ws_bom.cell(row=bom_row, column=10, value=bom_item.total_cost)
                ws_bom.cell(row=bom_row, column=11, value=bom_item.notes)
                ws_bom.cell(row=bom_row, column=12, value=bom_item.sequence)
                bom_row += 1

        # 第三个工作表：产品工序
        ws_process = wb.create_sheet(title="产品工序")
        
        # 工序表头
        process_headers = ['产品编码', '工序编号', '工序名称', '序号', '加工数量', '单价', '有效价格', '准备时间(分钟)', '加工时间(分钟)', '总时间(分钟)', '总成本', '是否必需', '备注']
        for col, header in enumerate(process_headers, 1):
            cell = ws_process.cell(row=1, column=col, value=header)
            ExcelGenerator._apply_header_style(cell)
        
        # 写入工序数据
        process_row = 2
        for product in products:
            for process_item in product.process_items.order_by('sequence'):
                ws_process.cell(row=process_row, column=1, value=product.product_code)
                ws_process.cell(row=process_row, column=2, value=process_item.process.process_code)
                ws_process.cell(row=process_row, column=3, value=process_item.process.process_name)
                ws_process.cell(row=process_row, column=4, value=process_item.sequence)
                ws_process.cell(row=process_row, column=5, value=process_item.quantity)
                ws_process.cell(row=process_row, column=6, value=process_item.unit_price)
                ws_process.cell(row=process_row, column=7, value=process_item.effective_price)
                ws_process.cell(row=process_row, column=8, value=process_item.setup_time)
                ws_process.cell(row=process_row, column=9, value=process_item.process_time)
                ws_process.cell(row=process_row, column=10, value=process_item.total_time)
                ws_process.cell(row=process_row, column=11, value=process_item.total_cost)
                ws_process.cell(row=process_row, column=12, value='是' if process_item.is_required else '否')
                ws_process.cell(row=process_row, column=13, value=process_item.notes)
                process_row += 1

        # 调整所有工作表的列宽
        for ws in [ws_product, ws_bom, ws_process]:
            for col in ws.columns:
                max_length = 0
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                ws.column_dimensions[col[0].column_letter].width = max_length + 2

        return wb 