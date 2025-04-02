import os
import sys
from datetime import datetime
from openpyxl import load_workbook

# 添加应用目录到路径
sys.path.append('.')

# 导入应用模块
from app.utils.excel_generator import ExcelGenerator

# 测试目录
TEST_DIR = 'test_data'
os.makedirs(TEST_DIR, exist_ok=True)

# 测试文件路径
TEMPLATE_PATH = os.path.join(TEST_DIR, 'test_template.xlsx')
TEST_PARSE_PATH = os.path.join(TEST_DIR, 'test_parse.xlsx')

def test_create_template():
    """测试模板创建功能"""
    print("测试1: 创建工序价格模板")
    try:
        wb = ExcelGenerator.create_process_price_template()
        wb.save(TEMPLATE_PATH)
        
        # 验证模板
        wb = load_workbook(TEMPLATE_PATH)
        ws = wb.active
        
        # 检查表头
        headers = [cell.value for cell in ws[1]]
        needed_headers = ['工序编号*', '工序名称*', '价格类型*', '包含工序']
        for header in needed_headers:
            if header not in headers:
                print(f"错误: 模板中缺少 '{header}' 列")
                return False
        
        print("模板创建成功，表头包含所有所需字段")
        
        # 检查示例数据
        for row in ws.iter_rows(min_row=2, max_row=5):
            row_values = [cell.value for cell in row]
            if not any(row_values):  # 跳过空行
                continue
            
            # 查找小计类型的行
            price_type_index = headers.index('价格类型*')
            included_processes_index = headers.index('包含工序')
            
            if row_values[price_type_index] == 'subtotal' and row_values[included_processes_index]:
                print("模板中包含小计示例数据，并且有包含工序")
                return True
        
        print("警告: 模板中找不到小计示例或包含工序")
        return False
    
    except Exception as e:
        print(f"创建模板失败: {str(e)}")
        return False

def test_parse_data():
    """测试数据解析功能"""
    print("\n测试2: 解析工序价格数据")
    try:
        # 创建测试数据文件
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        
        # 设置表头
        headers = ['工序编号*', '工序名称*', '部件', '图号', '型号', '单价*', '生效日期*', '备注', '价格类型*', '包含工序']
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        
        # 添加普通工序数据
        ws.cell(row=2, column=1, value='TEST001')
        ws.cell(row=2, column=2, value='测试工序1')
        ws.cell(row=2, column=3, value='测试部件')
        ws.cell(row=2, column=4, value='图号A')
        ws.cell(row=2, column=5, value='型号A')
        ws.cell(row=2, column=6, value=10.0)
        ws.cell(row=2, column=7, value=datetime.now().strftime('%Y-%m-%d'))
        ws.cell(row=2, column=8, value='测试备注')
        ws.cell(row=2, column=9, value='normal')
        
        # 添加第二个普通工序
        ws.cell(row=3, column=1, value='TEST002')
        ws.cell(row=3, column=2, value='测试工序2')
        ws.cell(row=3, column=3, value='测试部件')
        ws.cell(row=3, column=4, value='图号B')
        ws.cell(row=3, column=5, value='型号B')
        ws.cell(row=3, column=6, value=15.0)
        ws.cell(row=3, column=7, value=datetime.now().strftime('%Y-%m-%d'))
        ws.cell(row=3, column=8, value='测试备注')
        ws.cell(row=3, column=9, value='normal')
        
        # 添加小计工序
        ws.cell(row=4, column=1, value='TEST003')
        ws.cell(row=4, column=2, value='测试小计')
        ws.cell(row=4, column=3, value='测试部件')
        ws.cell(row=4, column=4, value='图号C')
        ws.cell(row=4, column=5, value='型号C')
        ws.cell(row=4, column=6, value='')  # 小计价格会自动计算
        ws.cell(row=4, column=7, value=datetime.now().strftime('%Y-%m-%d'))
        ws.cell(row=4, column=8, value='小计测试')
        ws.cell(row=4, column=9, value='subtotal')
        ws.cell(row=4, column=10, value='TEST001,TEST002')
        
        wb.save(TEST_PARSE_PATH)
        
        # 使用ExcelGenerator解析数据
        parsed_data = ExcelGenerator.parse_process_price_data(TEST_PARSE_PATH)
        if not parsed_data:
            print("错误: 解析数据为空")
            return False
        
        # 检查解析结果
        normal_count = 0
        subtotal_count = 0
        included_processes_found = False
        
        for data in parsed_data:
            if data['price_type'] == 'normal':
                normal_count += 1
            elif data['price_type'] == 'subtotal':
                subtotal_count += 1
                if data['included_processes'] and len(data['included_processes']) > 0:
                    included_processes_found = True
        
        print(f"解析结果: 普通工序: {normal_count}, 小计工序: {subtotal_count}")
        
        if normal_count == 2 and subtotal_count == 1 and included_processes_found:
            print("数据解析成功，可以识别普通工序和小计工序以及包含工序")
            return True
        else:
            print("错误: 数据解析不完整或不正确")
            return False
    
    except Exception as e:
        print(f"解析数据失败: {str(e)}")
        return False

def run_tests():
    """运行所有测试"""
    results = []
    
    # 测试1: 创建模板
    results.append(test_create_template())
    
    # 测试2: 解析数据
    results.append(test_parse_data())
    
    # 汇总结果
    success_count = results.count(True)
    fail_count = results.count(False)
    
    print("\n测试结果汇总:")
    print(f"总测试数: {len(results)}")
    print(f"成功: {success_count}")
    print(f"失败: {fail_count}")
    
    if fail_count == 0:
        print("\n所有测试通过！小计功能在Excel模板和数据解析中正常工作 🎉")
    else:
        print("\n有测试失败，请检查日志")

if __name__ == '__main__':
    run_tests() 